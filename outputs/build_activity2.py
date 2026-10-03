from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
OUT = Path(__file__).resolve().parent / 'Grupo_A1_M1.docx'

claude = '''La inteligencia artificial generativa como apoyo al aprendizaje universitario

Introducción

La irrupción de las herramientas de inteligencia artificial generativa, capaces de producir texto, código, resúmenes y explicaciones a partir de instrucciones en lenguaje natural, ha transformado la forma en que muchos estudiantes universitarios abordan sus tareas académicas. Esta transformación genera entusiasmo, pero también preocupación en las instituciones de educación superior. La pregunta central no es si estas herramientas deben existir en el entorno universitario, pues ya forman parte de él, sino cómo utilizarlas de manera que fortalezcan el aprendizaje en lugar de sustituirlo. Este ensayo sostiene que la inteligencia artificial generativa puede ser un apoyo valioso para estudiar y escribir, siempre que se emplee con criterio, se verifique la información que produce y se asuma plena responsabilidad académica sobre el resultado final.

Desarrollo
Usos para estudiar y escribir

En el estudio, estas herramientas pueden funcionar como un tutor disponible en cualquier momento. Un estudiante puede pedir que se le explique un concepto complejo con otras palabras, que se le propongan ejemplos adicionales o que se le formulen preguntas de práctica para prepararse ante una evaluación. También pueden ayudar a organizar un plan de estudio, a comparar teorías o a identificar vacíos en la comprensión cuando el estudiante explica un tema y solicita retroalimentación.

En la escritura, su utilidad se concentra en las etapas de planeación y revisión. Pueden ayudar a generar ideas iniciales, a estructurar un esquema, a mejorar la claridad de un párrafo o a detectar problemas de redacción, ortografía y coherencia. En estos casos, la herramienta actúa como un interlocutor que ofrece opciones, mientras que las decisiones sobre el contenido y los argumentos siguen en manos de quien escribe.

Beneficios

Uno de los principales beneficios es la personalización. A diferencia de una clase magistral, una herramienta conversacional permite adaptar el ritmo y el nivel de la explicación a las necesidades de cada estudiante. Además, ofrece retroalimentación inmediata, lo que facilita corregir errores a tiempo. También puede reducir el tiempo dedicado a tareas mecánicas, como dar formato o reorganizar ideas, de modo que el estudiante dedique más esfuerzo al análisis y la comprensión. Para quienes escriben en un idioma que no dominan por completo o enfrentan dificultades de redacción, puede representar además un apoyo para expresarse con mayor claridad.

Limitaciones y riesgos

Estas herramientas presentan limitaciones importantes. Generan respuestas con apariencia de seguridad y buena redacción, aun cuando el contenido sea impreciso o incorrecto. Es un fenómeno ampliamente reconocido, a veces llamado «alucinación», en el que el sistema puede producir datos, citas o referencias que no existen. Asimismo, sus respuestas dependen de los datos con los que fueron entrenadas, por lo que pueden reflejar sesgos o carecer de información actualizada sobre temas recientes o muy especializados.

Existe también un riesgo pedagógico: si el estudiante delega en la herramienta el esfuerzo de razonar, resumir o redactar, puede obtener un producto aceptable sin haber desarrollado las competencias que la tarea buscaba formar. El aprendizaje exige un trabajo cognitivo activo, y evitarlo de forma sistemática debilita precisamente las habilidades que la universidad pretende construir, como el pensamiento crítico y la argumentación propia.

La importancia de verificar la información

Por lo anterior, la verificación es una práctica indispensable. Toda afirmación relevante generada por una herramienta debe contrastarse con fuentes confiables, como libros de texto, artículos revisados por pares, documentos oficiales o los materiales de la asignatura. Las referencias bibliográficas que proponga el sistema deben comprobarse de manera independiente, pues pueden ser inexactas o inexistentes. Verificar no es un paso opcional; es lo que distingue el uso crítico de la dependencia ingenua, y lo que permite que la herramienta sea un punto de partida para la investigación y no su conclusión.

Responsabilidad académica

La responsabilidad académica es el eje que articula todo lo anterior. Quien entrega un trabajo responde por su contenido, sin importar qué herramientas haya utilizado para elaborarlo. Esto implica respetar las normas de integridad de su institución y las indicaciones de cada docente, que pueden permitir, restringir o condicionar el uso de estas tecnologías. Implica también ser transparente cuando corresponda, declarando el apoyo recibido, y evitar presentar como propio un texto que no refleja el propio esfuerzo intelectual. Las instituciones, por su parte, tienen el reto de establecer lineamientos claros y de formar a sus estudiantes en un uso ético y competente de estas herramientas.

Conclusión

La inteligencia artificial generativa ofrece oportunidades reales para el aprendizaje universitario: explicaciones personalizadas, práctica guiada y apoyo en la planeación y revisión de textos. Sin embargo, sus errores potenciales y el riesgo de reemplazar el esfuerzo intelectual exigen un uso prudente. La verificación sistemática de la información y la asunción de la responsabilidad académica son condiciones para que esta tecnología sume valor. En última instancia, el criterio del estudiante sigue siendo insustituible: la herramienta puede acompañar el proceso, pero el aprendizaje, y la responsabilidad por él, continúan siendo de quien estudia.'''

our_essay = [
'''La inteligencia artificial generativa ya forma parte de muchas actividades académicas. Con una instrucción escrita, estas herramientas pueden producir explicaciones, resúmenes, ejemplos y borradores. En la universidad, su presencia abre oportunidades para estudiar de manera interactiva y revisar textos con rapidez. También plantea una pregunta importante: ¿cómo aprovecharlas sin delegar en ellas el aprendizaje? Su valor depende de que el estudiante las use como apoyo, examine sus respuestas y conserve la responsabilidad sobre el trabajo que entrega.''',
'''Para estudiar, una herramienta generativa puede reformular un concepto con lenguaje más sencillo, proponer ejemplos o crear preguntas de práctica. El estudiante puede pedir aclaraciones sucesivas y comparar distintas explicaciones hasta detectar qué parte del tema aún no comprende. También puede emplearla para ordenar apuntes en un esquema o preparar un plan de repaso. Estas funciones facilitan el acceso inicial a una materia, pero la explicación obtenida necesita cotejarse con los materiales del curso, pues puede omitir matices o presentar una simplificación inadecuada.''',
'''En la escritura, la IA puede colaborar durante la lluvia de ideas, la organización de un argumento y la revisión de gramática o claridad. Por ejemplo, puede sugerir una estructura para un ensayo o señalar que un párrafo repite una idea. El estudiante debe decidir si las sugerencias sirven, comprobar que el texto sostenga su postura y reescribirlo con comprensión. De este modo, la herramienta ofrece alternativas para revisar, mientras que la selección de argumentos y la responsabilidad por el contenido permanecen en manos de quien escribe.''',
'''La rapidez y la disponibilidad son ventajas prácticas, pero no garantizan aprendizaje ni exactitud. Los sistemas generativos pueden redactar con seguridad incluso cuando se equivocan, mezclar conceptos o proponer referencias que no existen. También pueden producir respuestas generales que no atienden una consigna concreta. Por eso, antes de incorporar una afirmación a una tarea, conviene contrastarla con fuentes académicas confiables y verificar por separado los datos bibliográficos. Si la respuesta no puede confirmarse, no debe tratarse como evidencia.''',
'''Un segundo riesgo aparece cuando el estudiante entrega directamente una respuesta generada. En ese caso puede completar la tarea sin practicar la lectura crítica, la argumentación o la síntesis que la actividad busca desarrollar. Una estrategia más formativa consiste en usar la respuesta como borrador: identificar sus ideas, buscar evidencia, corregir errores y explicar el tema con palabras propias. La herramienta puede ahorrar tiempo en tareas mecánicas, pero ese tiempo solo beneficia al aprendizaje si se destina a comprender y revisar.''',
'''El uso responsable también depende de las normas de cada curso e institución. Antes de usar IA, el estudiante debe revisar qué permiten las instrucciones y cómo declarar el apoyo recibido. Si la herramienta colaboró en una parte del proceso, la transparencia ayuda a describir su papel sin atribuirle decisiones que tomó la persona. El estudiante sigue respondiendo por la precisión, las fuentes y la originalidad del trabajo presentado.''',
'''En conclusión, la IA generativa puede apoyar el estudio y la escritura universitaria al ofrecer explicaciones adaptables, preguntas de práctica y sugerencias de revisión. Sus límites exigen comprobar la información y valorar cada propuesta con criterio. Cuando el estudiante participa activamente, edita el resultado y sigue las pautas académicas, la herramienta puede complementar su proceso. La comprensión y la responsabilidad final continúan dependiendo de quien aprende.'''
]

comparison = '''Ambos textos presentan una tesis compatible: la IA generativa puede apoyar el aprendizaje universitario si el estudiante verifica sus respuestas y conserva la responsabilidad académica. El texto de Claude desarrolla el tema con mayor amplitud y una estructura explícita por apartados. Explica con detalle los usos para estudiar y escribir, los beneficios, los riesgos, la verificación y la integridad académica. Su organización facilita localizar cada subtema y ofrece una cobertura más completa de la consigna.

La versión generada para esta actividad presenta una secuencia más condensada. Integra los mismos ejes en párrafos continuos y conecta cada uso con una recomendación de revisión. Su lenguaje es claro y evita extenderse en afirmaciones que no se documentan en la consigna. En ambos textos la gramática y la coherencia son sólidas; ambos podrían requerir contraste de sus afirmaciones generales con fuentes del curso antes de una entrega académica. Ninguno incluye referencias verificables, por lo que no se deben tratar como investigación documentada.

Para desarrollar las secciones siguientes seleccionamos el texto de Claude como base, porque ofrece mayor cobertura temática y apartados que facilitan convertir sus ideas en preguntas didácticas. La elección se basa en organización y amplitud, no en una comprobación externa de cada afirmación. El texto requiere revisión humana y contraste con las instrucciones de la asignatura antes de presentarlo como versión final.'''

critical_review = '''El texto generado presenta una tesis clara y mantiene una secuencia lógica: comienza con usos posibles, explica beneficios, examina riesgos y concluye con la responsabilidad académica. La redacción es gramaticalmente fluida y los ejemplos de estudio y escritura ayudan a relacionar el tema con actividades universitarias. La extensión del ensayo es de aproximadamente 541 palabras, dentro del rango solicitado.

La revisión editorial se concentró en conservar una formulación equilibrada y evitar que las ventajas sonaran como resultados garantizados. También se revisó que las recomendaciones sobre verificación y transparencia fueran consistentes con la tesis. Como el texto no incorpora fuentes ni evidencia externa, sus afirmaciones generales deben contrastarse con las lecturas del curso antes de presentarlo como trabajo final. La versión no incluye citas ni datos específicos que permitan comprobarse desde el propio ensayo.

La herramienta fue útil para producir un borrador organizado con rapidez, pero el resultado por sí solo no demuestra comprensión ni reemplaza la revisión del estudiante. La intervención humana debe decidir qué argumentos son pertinentes, comprobar su precisión y adaptar el vocabulario al curso. En esta actividad también se comparó el borrador con el texto de Claude, lo cual permitió reconocer diferencias de amplitud y estructura. La selección final y cualquier declaración sobre el uso de IA deben ajustarse a las pautas indicadas por el docente.'''

reflection = '''La preparación del cuestionario permitió convertir las ideas principales del ensayo en tareas de aprendizaje. Las preguntas de selección múltiple y verdadero o falso comprueban nociones básicas, como el hecho de que una respuesta fluida puede contener errores y que las reglas del curso deben revisarse. Las preguntas abiertas y de aplicación exigen que el estudiante explique cómo verificaría una afirmación o actuaría ante una referencia dudosa. Así, el instrumento combina reconocimiento de conceptos con transferencia a situaciones académicas.

La IA puede acelerar la propuesta inicial de preguntas, pero la revisión humana es necesaria para asegurar que cada ítem sea claro, pertinente y tenga una respuesta defendible. También conviene evitar distractores confusos y revisar que las preguntas abiertas no sean tan amplias que resulte imposible valorar la respuesta. Por eso el cuestionario incluye respuestas orientativas en los ítems cerrados y criterios generales para las respuestas desarrolladas.

Como herramienta de aprendizaje, el cuestionario puede servir para repasar, detectar dudas y practicar el juicio crítico antes de una evaluación. Su utilidad aumenta cuando quien responde explica por qué eligió una opción y consulta fuentes para resolver errores. La actividad muestra que generar preguntas no basta: el grupo debe comprobar su calidad didáctica y su relación con el tema. La reflexión describe el diseño del material elaborado, no una medición empírica de aprendizaje.'''

questions = [
('Selección múltiple', '¿Cuál es un uso formativo de una herramienta generativa durante el estudio?', ['A. Entregar sin leer el primer resultado.', 'B. Pedir una explicación y contrastar los puntos importantes con materiales confiables.', 'C. Sustituir todas las lecturas de la asignatura.', 'D. Aceptar automáticamente las referencias sugeridas.'], 'Respuesta: B.'),
('Verdadero o falso', 'Una respuesta redactada con claridad es necesariamente correcta.', [], 'Respuesta: Falso. La fluidez no demuestra exactitud.'),
('Abierta', 'Menciona dos tareas académicas en las que la IA generativa podría servir como apoyo. Explica cómo conservarías la responsabilidad por el trabajo.', [], 'Criterio orientativo: identifica dos usos pertinentes y explica la revisión o decisión que mantiene el estudiante.'),
('Selección múltiple', '¿Cuál es un riesgo de entregar sin revisión un texto generado por IA?', ['A. Puede contener errores o ideas que el estudiante no comprende.', 'B. Siempre será demasiado corto.', 'C. Garantiza que todas las citas existan.', 'D. Asegura que el estudiante aprenda el tema.'], 'Respuesta: A.'),
('Verdadero o falso', 'Antes de usar IA en una tarea, es recomendable revisar las instrucciones del docente y del curso.', [], 'Respuesta: Verdadero.'),
('Abierta', '¿Qué pasos seguirías para comprobar una afirmación importante incluida en una respuesta de IA?', [], 'Criterio orientativo: contrastarla con materiales académicos o del curso y verificar la fuente original.'),
('Caso de aplicación', 'La herramienta propone una referencia bibliográfica que no logras localizar en el catálogo de la biblioteca ni en una base académica. ¿Qué harías antes de citarla? Explica por qué.', [], 'Criterio orientativo: no citarla hasta confirmar que existe y que respalda la afirmación; buscar datos verificables en fuentes confiables.'),
]

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.75)
sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for s in ['Title','Heading 1','Heading 2']:
    styles[s].font.name = 'Aptos Display'
    styles[s].font.color.rgb = RGBColor(25, 49, 74)
styles['Title'].font.size = Pt(25)
styles['Heading 1'].font.size = Pt(17)
styles['Heading 2'].font.size = Pt(13)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Actividad 2\nInteligencia artificial y aprendizaje universitario')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Integrantes: [Nombres del grupo]    Curso: [Nombre del curso]    Fecha: [Fecha]')
doc.add_paragraph('Este documento reúne los ensayos generados con Claude Sonnet 5.5 y ChatGPT, la comparación entre ambos, la selección del texto base y un cuestionario con reflexión. La comparación valora organización, claridad y cobertura; no sustituye la verificación de afirmaciones con fuentes académicas ni la revisión de las normas del curso.')

doc.add_heading('Sección 1. Ensayo de Claude', level=1)
doc.add_paragraph('Herramienta reportada por el estudiante: Claude Sonnet 5.5. Texto recibido, reproducido a continuación.')
for block in claude.split('\n\n'):
    if block in ['Introducción','Desarrollo','Usos para estudiar y escribir','Beneficios','Limitaciones y riesgos','La importancia de verificar la información','Responsabilidad académica','Conclusión']:
        doc.add_heading(block, level=2)
    else:
        doc.add_paragraph(block)

doc.add_heading('Sección 1. Ensayo generado con ChatGPT', level=1)
doc.add_paragraph('La inteligencia artificial generativa como apoyo al aprendizaje universitario', style='Heading 2')
for para in our_essay: doc.add_paragraph(para)
doc.add_heading('Análisis crítico del texto generado', level=2)
for para in critical_review.split('\n\n'): doc.add_paragraph(para)

doc.add_heading('Sección 2. Comparación y selección', level=1)
for para in comparison.split('\n\n'): doc.add_paragraph(para)

doc.add_heading('Sección 3. Cuestionario', level=1)
doc.add_paragraph('Tema: uso crítico y responsable de la inteligencia artificial generativa en la educación universitaria.')
for i,(kind,prompt,options,answer) in enumerate(questions,1):
    p=doc.add_paragraph(); r=p.add_run(f'{i}. {kind}. '); r.bold=True
    p.add_run(prompt)
    for opt in options: doc.add_paragraph(opt, style='List Bullet')
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.2)
    r=p.add_run(answer); r.italic=True

doc.add_heading('Reflexión sobre la elaboración del cuestionario', level=2)
for para in reflection.split('\n\n'): doc.add_paragraph(para)

doc.add_heading('Revisión final antes de entregar', level=1)
for item in ['Añadir nombres del grupo, curso y fecha.', 'Contrastar afirmaciones académicas importantes con fuentes del curso o fuentes confiables.', 'Confirmar con el docente cómo declarar el uso de herramientas generativas.', 'Editar la redacción final para que refleje el trabajo y criterio del grupo.']:
    doc.add_paragraph(item, style='List Bullet')

# Footer with page number field
for section in doc.sections:
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run('Actividad 2  |  ')
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); fp._p.append(fld)

doc.save(OUT)
print(OUT)
