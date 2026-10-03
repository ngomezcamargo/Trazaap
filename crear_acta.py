from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = 'Acta_entrega_satisfaccion_Trazaap.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)
normal = doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(9)
normal.paragraph_format.space_after=Pt(4)

def shade(cell, color):
    pr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),color); pr.append(sh)
def set_cell(cell, text, bold=False, color=None):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(text); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(8.5)
    if color: r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('ACTA DE ENTREGA A SATISFACCIÓN Y CIERRE DEL PROYECTO'); r.bold=True; r.font.name='Arial'; r.font.size=Pt(13); r.font.color.rgb=RGBColor(31,61,87)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Proyecto de grado · Programa de Ingeniería de Sistemas · Universidad El Bosque'); r.italic=True; r.font.size=Pt(9)

for label, value in [
 ('Nombre completo del proyecto','Trazaap: [completar título oficial registrado en la propuesta]'),
 ('Autores (estudiantes)','[Nombre completo de cada autor]'),
 ('Código asignado por el Programa','[Código del proyecto]'),
 ('Organización beneficiaria','[Razón social]'),
 ('Fecha y lugar','[Ciudad], [día] de [mes] de [año]')]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    r=p.add_run(label+': '); r.bold=True; p.add_run(value)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5)
p.add_run('Declaración de cierre. ').bold=True
p.add_run('Mediante la presente acta se deja constancia del cierre del proyecto de grado indicado, desarrollado como ejercicio académico en el marco del trabajo de grado del Programa de Ingeniería de Sistemas de la Universidad El Bosque.')

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(4); p.add_run('Relación de entregables').bold=True
p=doc.add_paragraph('El representante del beneficiario verifica la recepción de los siguientes productos acordados para el proyecto:'); p.paragraph_format.space_after=Pt(3)

rows=[
 ('1','Aplicación web Trazaap para gestión de proveedores, materias primas, recepción e inspección, producción, liberación, inventario, trazabilidad y reportes.','[Repositorio / URL / medio de entrega]','☐ Recibido'),
 ('2','Backend y API del sistema, con persistencia de datos y servicios de negocio.','[Repositorio / medio de entrega]','☐ Recibido'),
 ('3','Componente de trazabilidad verificable con Hyperledger Fabric y chaincode asociado.','[Repositorio / medio de entrega]','☐ Recibido'),
 ('4','Documentación técnica y funcional del sistema.','[Indicar documentos entregados]','☐ Recibido'),
 ('5','Plan y resultados de pruebas del sistema.','[Indicar documentos/evidencias entregados]','☐ Recibido'),
]
t=doc.add_table(rows=1, cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; t.autofit=False
widths=[.38,3.15,2.15,.85]
for c,w,txt in zip(t.rows[0].cells,widths,['N.º','Entregable','Medio o evidencia de entrega','Verificación']):
    c.width= Inches(w); set_cell(c,txt,True,'FFFFFF'); shade(c,'1F3D57')
for row in rows:
    cells=t.add_row().cells
    for c,w,txt in zip(cells,widths,row): c.width= Inches(w); set_cell(c,txt)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(4)
p.add_run('Constancia de recibo. ').bold=True
p.add_run('El(La) representante del beneficiario declara que recibió a satisfacción la totalidad de los entregables relacionados en esta acta, de acuerdo con lo pactado para el proyecto. En consecuencia, certifica que el proyecto queda terminado y cerrado, sin entregables pendientes.')

p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
p.add_run('Anexo obligatorio: ').bold=True; p.add_run('Formato de licencia de autorización de uso diligenciado en Proyecto 1 y firmado por el beneficiario.')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
p.add_run('Identificación formal: ').bold=True; p.add_run('Esta acta se suscribe en papel membreteado de la organización beneficiaria. Si la organización no cuenta con correo institucional o elementos distintivos, se adjunta certificado de existencia y representación legal u otro documento que la identifique formalmente.')

doc.add_paragraph('\n')
t2=doc.add_table(rows=1, cols=2); t2.alignment=WD_TABLE_ALIGNMENT.CENTER; t2.autofit=False
for c in t2.rows[0].cells: c.width= Inches(3.1)
set_cell(t2.cell(0,0),'__________________________________\nFirma del representante del beneficiario\nNombre: [nombre completo]\nCargo: [cargo]\nCorreo institucional/corporativo: [correo]\nTeléfono: [teléfono]')
set_cell(t2.cell(0,1),'__________________________________\nNombre y firma de representante(s) del proyecto\nNombre(s): [nombre(s) completo(s)]')

doc.core_properties.title='Acta de entrega a satisfacción y cierre del proyecto Trazaap'
doc.core_properties.subject='Proyecto de grado - Universidad El Bosque'
doc.save(out)
print(out)
