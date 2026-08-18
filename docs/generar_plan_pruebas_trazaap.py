from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "docs" / "Plan_de_Pruebas_Trazaap.docx"
LOGO = BASE / "frontend" / "public" / "trazaap-logo.jpeg"

GREEN = "0B6B35"
DARK_GREEN = "064E2A"
LIGHT_GREEN = "E8F5EC"
PALE_GREEN = "F4FAF6"
GOLD = "A56A00"
LIGHT_GOLD = "FFF4D6"
BLUE = "245B78"
LIGHT_BLUE = "EAF3F8"
RED = "A61B1B"
LIGHT_RED = "FDECEC"
GRAY = "667085"
LIGHT_GRAY = "F3F5F7"
TEXT = "17211B"
WHITE = "FFFFFF"
BORDER = "C8D8CE"

STATUS_AUTO = "Aprobada automatizada"
STATUS_EVIDENCE = "Pendiente de evidencia"
STATUS_EXECUTE = "Pendiente de ejecución"
STATUS_IMPLEMENT = "Pendiente de implementación"


def case(case_id, requirement, test, expected, justification, status, evidence="Por adjuntar"):
    return {
        "id": case_id,
        "req": requirement,
        "test": test,
        "expected": expected,
        "why": justification,
        "status": status,
        "evidence": evidence,
    }


functional_cases = [
    case("PF-001", "RF12", "Inicio de sesión válido", "El usuario activo entra con su perfil y llega al panel autorizado.", "Confirma que la identidad se valida antes de exponer información operativa.", STATUS_EVIDENCE),
    case("PF-002", "RF12", "Credenciales incorrectas o usuario inactivo", "La API rechaza el acceso sin revelar si falló el correo, la contraseña o el estado de la cuenta.", "Evita accesos indebidos y reduce la enumeración de usuarios.", STATUS_EVIDENCE),
    case("PF-003", "RF12 / RNF01", "Credenciales fuera de la URL", "Correo y contraseña viajan en el cuerpo POST y no aparecen en historial, parámetros ni enlaces.", "Las credenciales visibles en la URL pueden quedar en historial, logs y capturas.", STATUS_EVIDENCE),
    case("PF-004", "RF12", "Cierre manual y expiración por inactividad", "Cerrar sesión elimina la sesión; tras dos horas sin actividad se exige autenticarse nuevamente.", "Limita el uso de una sesión abandonada en equipos compartidos de planta.", STATUS_EVIDENCE),
    case("PF-005", "RF12", "Administrador en ruta restringida", "El middleware permite la operación administrativa.", "Demuestra que el rol con mayor autoridad conserva las funciones de gestión.", STATUS_AUTO, "backend/test/rbac.middleware.test.js"),
    case("PF-006", "RF12", "Rol expresamente autorizado", "La ruta permite al rol declarado en su política.", "Verifica que la autorización depende de una regla explícita y no solo de estar autenticado.", STATUS_AUTO, "backend/test/rbac.middleware.test.js"),
    case("PF-007", "RF12", "Rol sin permiso", "La API responde 403 y no ejecuta cambios.", "El frontend puede ocultar botones, pero la protección real debe existir en el backend.", STATUS_AUTO, "backend/test/rbac.middleware.test.js"),
    case("PF-008", "RF12", "Rol inválido o desconocido", "La API responde 403 y no asigna permisos por defecto.", "Impide que errores de configuración se conviertan en privilegios involuntarios.", STATUS_AUTO, "backend/test/rbac.middleware.test.js"),
    case("PF-009", "RF10-RF12", "Códigos externos estables", "El mismo lote y contexto producen códigos consistentes para cliente y auditoría.", "El enlace de verificación debe seguir funcionando después de generarse o imprimirse.", STATUS_AUTO, "backend/test/codigos-acceso.test.js"),
    case("PF-010", "RF10-RF12", "Normalización de código externo", "Variaciones permitidas de escritura se comparan sin perder seguridad.", "Reduce falsos rechazos al transcribir un código desde una factura o etiqueta.", STATUS_AUTO, "backend/test/codigos-acceso.test.js"),
    case("PF-011", "RF12 / RNF01", "Token ausente, vencido o alterado", "La API responde 401 y no retorna ni modifica datos.", "Comprueba que un token inválido no sea tratado como una sesión válida.", STATUS_EXECUTE),
    case("PF-012", "RF12", "Gestión de usuarios", "El administrador crea, consulta, actualiza, cambia rol y desactiva cuentas sin duplicar correos.", "Los permisos dependen de cuentas correctamente administradas durante todo su ciclo de vida.", STATUS_EVIDENCE),
    case("PF-013", "RF17", "Crear proveedor", "Se guardan razón social, NIT, contacto, teléfono, correo, dirección y certificaciones.", "La trazabilidad hacia atrás necesita identificar el origen comercial de cada insumo.", STATUS_EVIDENCE),
    case("PF-014", "RF17", "Actualizar, desactivar y evitar duplicados de proveedor", "Los cambios válidos persisten; NIT duplicado o proveedor inactivo se controla según la regla definida.", "Evita inconsistencias en recepciones y conserva el historial sin borrar proveedores usados.", STATUS_EVIDENCE),
    case("PF-015", "RF01-A / RF17", "Crear y actualizar materia prima", "La materia prima conserva nombre, descripción, tipo, unidad base y almacenamiento.", "La receta, el inventario y la recepción comparten este catálogo maestro.", STATUS_EVIDENCE),
    case("PF-016", "RF01-A", "Presentaciones y unidades de materia prima", "Las presentaciones se registran con cantidades válidas y se convierten a la unidad base correcta.", "Un error de unidad altera inventario, receta y consumo de producción.", STATUS_EVIDENCE),
    case("PF-017", "RF17", "Relación proveedor-materia prima y certificaciones", "La asociación se consulta desde recepciones y trazabilidad sin perder la ficha del proveedor.", "Permite demostrar procedencia y condición sanitaria del suministro.", STATUS_EVIDENCE),
    case("PF-018", "RF01-A / RF17", "Catálogos activos en formularios", "Los formularios muestran únicamente opciones utilizables y conservan referencias históricas inactivas.", "Evita seleccionar registros deshabilitados sin destruir la trazabilidad existente.", STATUS_EVIDENCE),
    case("PF-019", "RF01-A", "Registrar recepción válida", "La recepción queda vinculada con proveedor, materia prima, lotes, cantidades, fechas y responsable.", "Es el punto inicial de la trazabilidad y debe capturar qué llegó, de quién y cuándo.", STATUS_EVIDENCE),
    case("PF-020", "RF01-A / RNF06", "Validaciones negativas de recepción", "Campos obligatorios vacíos, cantidades no positivas y fechas incoherentes se rechazan dentro del formulario.", "Previene que datos incompletos lleguen a inventario y blockchain.", STATUS_EVIDENCE),
    case("PF-021", "RF01-A", "Ingreso de recepción al inventario", "Una recepción aceptada aumenta exactamente el stock de la materia prima en su unidad base.", "El inventario debe representar la disponibilidad física derivada de recepciones reales.", STATUS_EVIDENCE),
    case("PF-022", "RF01-A", "Movimiento de inventario por recepción", "Se registra tipo, cantidad, unidad, referencia, proveedor, responsable y fecha.", "El saldo por sí solo no explica su origen; el movimiento permite auditarlo.", STATUS_EVIDENCE),
    case("PF-023", "RF01-B", "Inspección organoléptica y de empaque", "Olor, color, textura, empaque, certificado y observaciones quedan ligados a la recepción.", "Documenta la evaluación sanitaria que sustenta la aceptación del insumo.", STATUS_EVIDENCE),
    case("PF-024", "RF01-B", "Inspección de transporte", "Condiciones del vehículo e higiene del conductor se registran como cumple/no cumple.", "Las condiciones de transporte pueden afectar la inocuidad antes del ingreso a planta.", STATUS_EVIDENCE),
    case("PF-025", "RF01-B", "Decisión de aceptación", "Una inspección conforme deja la recepción disponible y registra la decisión y responsable.", "La materia prima no debe utilizarse sin una decisión explícita de recepción.", STATUS_EVIDENCE),
    case("PF-026", "RF01-B", "Decisión de rechazo", "El rechazo conserva la evidencia y evita tratar el insumo como disponible.", "Impide consumir materia prima no conforme y conserva el soporte de la decisión.", STATUS_EVIDENCE),
    case("PF-027", "RF01-B", "Inspección única por recepción", "El sistema evita duplicar la inspección o aplica una regla de corrección controlada.", "Dos decisiones finales sobre la misma recepción generan ambigüedad operativa.", STATUS_EXECUTE),
    case("PF-028", "RF02", "Configurar producto y estándares", "Se guardan categoría, vida útil, almacenamiento, tiempos y temperaturas esperadas.", "Los estándares permiten comparar la planeación con lo ocurrido en manufactura.", STATUS_EVIDENCE),
    case("PF-029", "RF02", "Variantes mínimas y eliminación de adicionales", "Existe al menos una presentación; las adicionales pueden quitarse con la X.", "Evita productos sin presentación y corrige adiciones accidentales antes de guardar.", STATUS_EVIDENCE),
    case("PF-030", "RF02 / RNF06", "Cancelar y reiniciar formularios de producto", "Al cancelar se borran los valores y al reabrir aparece una única variante vacía.", "Datos residuales pueden crear productos o recetas equivocados.", STATUS_EVIDENCE),
    case("PF-031", "RF02", "Receta dependiente de tamaño", "Cada variante conserva sus materias primas y cantidades por unidad.", "Una presentación grande y una pequeña no deben descontar la misma cantidad de insumos.", STATUS_EVIDENCE),
    case("PF-032", "RF02 / RF07", "Inmersión exclusiva para bagels", "Los campos de inmersión aparecen y se validan en bagels; en otros panes no se exigen.", "La captura debe representar el proceso real y evitar controles irrelevantes.", STATUS_EVIDENCE),
    case("PF-033", "RF02", "Conversión y edición de receta", "La receta editada mantiene unidades compatibles y recalcula necesidades sin duplicar ingredientes.", "Protege el cálculo de abastecimiento frente a cambios de producto.", STATUS_EVIDENCE),
    case("PF-034", "RF02", "Crear orden con un producto", "La orden se crea sin responsable individual, con fecha, producto, variante y cantidad programada.", "La orden planifica el trabajo colectivo y no debe confundirse con la ejecución real.", STATUS_EVIDENCE),
    case("PF-035", "RF02 / RNF06", "Agregar y retirar productos de la orden", "Los productos adicionales pueden eliminarse; cancelar limpia todo y deja una fila vacía.", "Permite corregir errores de digitación sin abandonar o dañar la orden.", STATUS_EVIDENCE),
    case("PF-036", "RF02", "Cálculo automático de insumos", "Cantidad por unidad multiplicada por cantidad programada produce el total correcto por materia prima.", "El cálculo evita planeación manual inconsistente y anticipa necesidades de inventario.", STATUS_EVIDENCE),
    case("PF-037", "RF02", "Inventario insuficiente", "La orden informa el faltante y aplica la regla definida sin confundir referencias estándar con insumos.", "No se debe iniciar una producción que físicamente no puede abastecerse.", STATUS_EVIDENCE),
    case("PF-038", "RF02", "Persistencia y consulta de orden", "La orden nueva aparece en el listado y en orden activa con productos y materias correctos.", "Confirma que la transacción completa se guardó y puede continuar su flujo.", STATUS_EVIDENCE),
    case("PF-039", "RF02", "Transiciones de estado de orden", "Solo se permiten cambios coherentes entre pendiente, en proceso, lista para liberación y finalizada.", "Los estados controlan qué acciones posteriores son válidas.", STATUS_EVIDENCE),
    case("PF-040", "RF02", "Manufactura asociada a una orden", "No se registra manufactura sin orden y producto programado existentes.", "Mantiene la relación entre lo planeado y lo realmente producido.", STATUS_EVIDENCE),
    case("PF-041", "RF02 / RF12", "Responsable de manufactura operario", "El selector muestra usuarios activos con rol operario y guarda su nombre como responsable.", "La responsabilidad debe corresponder a una cuenta identificable y autorizada.", STATUS_EVIDENCE),
    case("PF-042", "RF02 / RF07", "Valores reales frente a estándares", "Tiempos y temperaturas reales se guardan y las desviaciones generan advertencia sin bloquear.", "Permite detectar variaciones del proceso y conservar la realidad operativa.", STATUS_EVIDENCE),
    case("PF-043", "RF02", "Campos reales de inmersión", "Son obligatorios solo cuando el producto requiere inmersión.", "Evita registros incompletos de bagels y datos ficticios en otros productos.", STATUS_EVIDENCE),
    case("PF-044", "RF02", "Validar lote, unidades y horario de manufactura", "Lote no vacío, unidades positivas y hora fin posterior a hora inicio.", "Impide eventos productivos imposibles o imposibles de rastrear.", STATUS_EVIDENCE),
    case("PF-045", "RF02 / RNF06", "Confirmación y cierre de manufactura", "Al guardar correctamente se cierra el modal y aparece una confirmación visible; el doble envío queda bloqueado.", "El operario debe saber inequívocamente si el registro fue aceptado.", STATUS_EVIDENCE),
    case("PF-046", "RF02", "Finalización automática de orden", "Cuando se completan las manufacturas programadas la orden pasa a finalizada.", "La ejecución completa debe reflejarse sin depender de una acción manual redundante.", STATUS_EVIDENCE),
    case("PF-047", "RF02", "Consumo real de materias primas", "Manufactura descuenta cantidades reales, selecciona lotes disponibles y genera movimientos auditables.", "Relaciona físicamente las materias primas de origen con el lote terminado.", STATUS_EVIDENCE),
    case("PF-048", "RF03-A", "Productos pendientes de liberación", "Solo manufacturas registradas y no liberadas aparecen como pendientes.", "Evita liberar un producto inexistente o procesarlo dos veces.", STATUS_EVIDENCE),
    case("PF-049", "RF03-A / RF08", "Validaciones obligatorias de liberación", "Etiqueta y envase deben marcarse; los errores aparecen dentro del modal.", "La salida requiere evidencia de que el producto terminado fue revisado.", STATUS_EVIDENCE),
    case("PF-050", "RF03-A", "Liberación aprobada", "Crea inventario terminado con lote, unidades, vencimiento y estado disponible.", "Solo el producto aprobado puede quedar disponible para salida.", STATUS_EVIDENCE),
    case("PF-051", "RF03-A / RF08", "Liberación retenida", "Exige motivo y no habilita el producto como disponible.", "La retención separa temporalmente un producto que requiere decisión posterior.", STATUS_EVIDENCE),
    case("PF-052", "RF03-A / RF08", "Liberación rechazada", "Exige motivo, conserva el evento y bloquea la salida.", "Un producto no conforme no puede mezclarse con inventario liberado.", STATUS_EVIDENCE),
    case("PF-053", "RF03-A", "Evitar doble liberación", "Una manufactura o lote ya liberado no puede liberarse nuevamente.", "Previene duplicar inventario terminado y decisiones contradictorias.", STATUS_EVIDENCE),
    case("PF-054", "RF03-A / RF05", "Responsable, factura y transporte", "Se guardan operario, número de factura, conductor, placa, limpieza y documentación.", "Estos datos identifican quién autorizó y bajo qué condiciones salió el producto.", STATUS_EVIDENCE),
    case("PF-055", "RF03-A", "Cantidades y vencimiento de liberación", "Unidades positivas, no superiores a las producidas, peso válido y vencimiento posterior.", "Evita liberar cantidades inexistentes o productos vencidos.", STATUS_EXECUTE),
    case("PF-056", "RF03-A / RNF06", "Confirmación y cierre de liberación", "Durante el guardado se muestra progreso; al éxito se cierra el modal y queda una confirmación visible.", "Reduce dobles envíos y elimina la incertidumbre del usuario.", STATUS_EVIDENCE),
    case("PF-057", "RF05", "Salida y descuento de producto terminado", "La salida reduce exactamente el lote liberado y conserva factura, conductor y movimiento.", "El inventario final debe coincidir con las unidades que permanecen en planta.", STATUS_EVIDENCE),
    case("PF-058", "RF06", "Consulta por lote final", "El sistema reconstruye recepción, inspección, orden, manufactura, liberación e inventario relacionados.", "El lote producido es el identificador que acompaña al alimento hacia el cliente.", STATUS_EVIDENCE),
    case("PF-059", "RF06", "Trazabilidad hacia atrás", "Desde el producto se identifican materias primas, lotes recibidos y proveedores.", "Permite localizar el origen ante una no conformidad o retiro.", STATUS_EVIDENCE),
    case("PF-060", "RF05-RF06", "Trazabilidad hacia adelante hasta cliente", "Desde una materia prima se obtienen lotes producidos, salidas y clientes receptores.", "Es necesaria para informar a todos los afectados por un insumo comprometido.", STATUS_IMPLEMENT),
    case("PF-061", "RF06 / RF09", "Estados de validación en trazabilidad", "Cada evento muestra verificado, alterado, no encontrado o pendiente según Fabric.", "Diferencia la existencia del dato operativo de su evidencia criptográfica.", STATUS_EVIDENCE),
    case("PF-062", "RF10", "QR único y acceso público", "El QR abre el lote correcto sin autenticación y conserva legibilidad al imprimir.", "Vincula el producto físico con su historial digital.", STATUS_EVIDENCE),
    case("PF-063", "RF11 / RNF07", "Privacidad del portal público", "El consumidor ve procedencia y proceso sin datos internos sensibles; el auditor recibe información autorizada adicional.", "Cumple transparencia sin exponer información comercial o personal.", STATUS_EVIDENCE),
    case("PF-064", "RF14", "Reporte PDF de trazabilidad", "El PDF contiene eventos, responsables, validación Fabric y QR sin cortes ni solapamientos.", "Es la evidencia compartible para auditoría y consulta fuera de la aplicación.", STATUS_EVIDENCE),
    case("PF-065", "RF14", "Reporte Excel y consulta de hasta 50 lotes", "Se genera un archivo estructurado dentro del tiempo requerido y con estados blockchain.", "El requisito exige análisis masivo y exportación reutilizable, no solo impresión individual.", STATUS_IMPLEMENT),
    case("PF-066", "RF04", "Condiciones de almacenamiento", "Registrar ubicación, temperatura, ingreso, salida y refrigeración por lote.", "La trazabilidad incluye la etapa entre producción y distribución.", STATUS_IMPLEMENT),
    case("PF-067", "RF07", "Puntos críticos HACCP", "Registrar variable, rango, medición, desviación, alerta y decisión vinculada al lote.", "Los puntos críticos demuestran control preventivo del proceso.", STATUS_IMPLEMENT),
    case("PF-068", "RF08", "Controles de calidad e inocuidad", "Registrar resultados físicos, químicos, microbiológicos y organolépticos con decisión final.", "La liberación debe estar sustentada por controles de calidad trazables.", STATUS_IMPLEMENT),
    case("PF-069", "RF15", "Plan de saneamiento", "Registrar limpieza, desinfección, plagas, residuos, cronograma, responsable y resultado.", "La Resolución 2674 exige evidencia recuperable de las prácticas de saneamiento.", STATUS_IMPLEMENT),
    case("PF-070", "RF16", "Devoluciones y lotes no conformes", "Relacionar lote, cliente, motivo, decisión, responsable y destino del producto.", "Completa la trazabilidad posterior a la entrega y el manejo de incidentes.", STATUS_IMPLEMENT),
    case("PF-071", "RF17", "Documentos de proveedor en almacenamiento de objetos", "Adjuntar certificado o ficha, recuperar el archivo y verificar su referencia o huella.", "Las evidencias documentales no deben quedar dispersas ni sobrecargar PostgreSQL/Fabric.", STATUS_IMPLEMENT),
]


blockchain_cases = [
    case("BC-001", "RF09", "Hash estable ante distinto orden de claves", "Dos objetos funcionalmente iguales producen el mismo SHA-256.", "La validación no debe depender del orden accidental de serialización.", STATUS_AUTO, "backend/test/trazabilidad.hash.test.js"),
    case("BC-002", "RF01-A / RF09", "Payload estable de recepción", "La normalización incluye datos funcionales y excluye metadatos técnicos.", "El hash debe representar el evento de recepción y poder recalcularse.", STATUS_AUTO, "backend/test/trazabilidad.hash.test.js"),
    case("BC-003", "RF01-B / RF09", "Payload estable de inspección", "Incluye decisión, vehículo y conductor con estructura determinista.", "Una modificación de la inspección debe detectarse sin falsos cambios técnicos.", STATUS_AUTO, "backend/test/trazabilidad.hash.test.js"),
    case("BC-004", "RF13", "Fabric no disponible durante despacho", "La API responde 503 y no escribe el despacho.", "Evita afirmar una salida validada cuando la regla crítica no pudo ejecutarse.", STATUS_AUTO, "backend/test/rf13-chaincode-backend.test.js"),
    case("BC-005", "RF13", "Despacho bloqueado sin efectos operativos", "No descuenta inventario ni crea liberación.", "Una regla inteligente debe ser atómica con la operación que controla.", STATUS_AUTO, "backend/test/rf13-chaincode-backend.test.js"),
    case("BC-006", "RF13", "Error de registro revierte transacción", "PostgreSQL vuelve al estado anterior cuando falla el registro crítico.", "Previene divergencia entre el estado operativo y la evidencia blockchain.", STATUS_AUTO, "backend/test/rf13-chaincode-backend.test.js"),
    case("BC-007", "RF11 / RF13", "Confirmación pública con credenciales válidas", "Solo un código externo válido envía la confirmación a Fabric.", "El portal público no debe aceptar confirmaciones anónimas o falsificadas.", STATUS_AUTO, "backend/test/rf13-chaincode-backend.test.js"),
    case("BC-008", "RF09", "Registrar evento nuevo", "El chaincode guarda evento, actor Fabric, lote, payload, hash y timestamp.", "Demuestra la creación real de evidencia inmutable en el ledger.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-009", "RF09", "Rechazar sobrescritura", "La clave existente no se reemplaza silenciosamente.", "La inmutabilidad lógica exige impedir que el último valor borre la evidencia anterior.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-010", "RF06 / RF09", "Detectar registro alterado", "El hash recalculado distinto retorna ALTERADO.", "Es la prueba central de integridad prometida por la arquitectura.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-011", "RF09", "Corrección sin modificar original", "La corrección crea una nueva evidencia y conserva el evento previo.", "Corregir datos no debe destruir el historial de auditoría.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-012", "RF05 / RF13", "Confirmar recepción por cliente", "Un despacho aprobado recibe una confirmación autenticada.", "Cierra la trazabilidad hacia adelante con evidencia del receptor.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-013", "RF13", "Impedir doble confirmación", "Una segunda confirmación del mismo despacho se rechaza.", "Evita estados contradictorios y transacciones redundantes.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-014", "RF05 / RF13", "Despacho conforme", "Las reglas aceptan y registran un despacho dentro de parámetros.", "Demuestra que el contrato no solo bloquea, también habilita operaciones válidas.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-015", "RF07 / RF13", "Bloqueo por control crítico", "Una variable fuera de rango impide el despacho y retorna motivos.", "Automatiza un control preventivo que no debe depender solo de revisión humana.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-016", "RF13", "Bloqueo de lote vencido", "El chaincode rechaza la salida cuando la fecha de vencimiento fue superada.", "Impide entregar producto que ya no es apto según su vida útil registrada.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
    case("BC-017", "RF13", "Alerta única por vencimiento", "Se registra una alerta por lote sin duplicarla en consultas repetidas.", "La idempotencia evita ruido operativo y múltiples evidencias para el mismo hecho.", STATUS_AUTO, "fabric/chaincode/traceability/test/traceability.test.js"),
]


nonfunctional_cases = [
    case("PNF-001", "RNF01", "Seguridad e integridad integral", "SQL injection, XSS, escalamiento de rol y manipulación de JWT no comprometen datos ni permisos.", "La integridad blockchain no reemplaza la seguridad de aplicación y base de datos.", STATUS_EXECUTE),
    case("PNF-002", "RNF02", "Disponibilidad mensual", "Monitoreo demuestra al menos 99 % fuera de mantenimientos notificados.", "El porcentaje requiere evidencia temporal; no puede concluirse con una demostración local.", STATUS_EXECUTE),
    case("PNF-003", "RNF03", "Secuencia temporal verificable", "Fechas operativas y timestamps Fabric reconstruyen el orden exacto de cada lote.", "Las auditorías deben demostrar que los controles ocurrieron en el momento correcto.", STATUS_EXECUTE),
    case("PNF-004", "RNF04", "Matriz Resolución 2674", "Cada campo normativo se relaciona con formulario, tabla, endpoint, RF y evidencia de prueba.", "Permite sustentar cumplimiento sin depender de afirmaciones generales.", STATUS_EXECUTE),
    case("PNF-005", "RNF05", "Tiempo de consulta", "Trazabilidad y operaciones habituales responden en menos de tres segundos bajo carga normal.", "Una consulta lenta limita la reacción ante auditoría o retiro de producto.", STATUS_EXECUTE),
    case("PNF-006", "RNF05", "Tiempo de confirmación blockchain", "Un evento se confirma o informa su fallo en máximo diez segundos.", "El operario necesita retroalimentación rápida para no repetir registros.", STATUS_EXECUTE),
    case("PNF-007", "RNF06", "Usabilidad controlada", "Un operario sin capacitación previa completa manufactura en menos de tres minutos y entiende los errores.", "La adopción depende de usuarios con alfabetización digital básica.", STATUS_EXECUTE),
    case("PNF-008", "RNF07", "Confidencialidad diferenciada", "Consumidor, auditor y usuario interno reciben únicamente los campos autorizados.", "La transparencia pública no debe revelar condiciones comerciales o datos personales.", STATUS_EXECUTE),
    case("PNF-009", "RNF08", "Fallo de nodos y consenso", "La red mantiene o rechaza operaciones de forma predecible al caer un peer u orderer y recupera el estado.", "La tolerancia a fallos solo se demuestra provocando fallas reales en la topología.", STATUS_EXECUTE),
    case("PNF-010", "RNF09", "Escalabilidad a diez veces el volumen", "Con diez veces lotes y eventos se mantienen los límites de respuesta sin rediseño funcional.", "Comprueba que el crecimiento no degrada la utilidad del sistema.", STATUS_EXECUTE),
    case("PNF-011", "RNF10", "Mantenibilidad y regresión", "Builds, migraciones y pruebas pasan después de cambios; código, contratos y arquitectura están documentados.", "Reduce el riesgo de que una corrección rompa módulos ya terminados.", STATUS_EXECUTE),
    case("PNF-012", "RNF11", "Interoperabilidad", "La API entrega JSON documentado y se valida la brecha frente a OAuth 2.0 y GS1 EPCIS 2.0.", "Permite integrar inventario, logística, calidad y otros actores sin acoplamiento propietario.", STATUS_IMPLEMENT),
    case("PNF-013", "Transversal", "Compatibilidad y accesibilidad", "Flujos críticos funcionan en Chrome, Edge, Firefox y móvil, con teclado, etiquetas y contraste legibles.", "La aplicación debe ser operable desde los dispositivos reales de planta y consulta pública.", STATUS_EXECUTE),
]


ALL_CASES = functional_cases + blockchain_cases + nonfunctional_cases


def set_run(run, size=10, bold=False, color=TEXT, font="Aptos", italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=0):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, width in enumerate(widths):
            set_cell_width(row.cells[index], width)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, size=8.4, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text))
    set_run(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def status_fill(status):
    if status == STATUS_AUTO:
        return LIGHT_GREEN, DARK_GREEN
    if status == STATUS_EVIDENCE:
        return LIGHT_GOLD, GOLD
    if status == STATUS_EXECUTE:
        return LIGHT_BLUE, BLUE
    return LIGHT_RED, RED


def add_page_number(paragraph):
    run = paragraph.add_run("Página ")
    set_run(run, size=8, color=GRAY)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def configure_header_footer(section, label="PLAN DE PRUEBAS | TRAZAAP", different_first=False):
    section.different_first_page_header_footer = different_first
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    set_run(paragraph.add_run(label), size=8, bold=True, color=GREEN)
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    add_page_number(paragraph)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    heading_tokens = {
        "Heading 1": (16, GREEN, 16, 7),
        "Heading 2": (13, GREEN, 12, 5),
        "Heading 3": (11, DARK_GREEN, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        set_run(paragraph.add_run(bold_lead), size=10, bold=True)
        set_run(paragraph.add_run(text[len(bold_lead):]), size=10)
    else:
        set_run(paragraph.add_run(text), size=10)
    return paragraph


def add_callout(doc, title, body, fill=PALE_GREEN, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [9360], indent=0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    set_run(paragraph.add_run(title), size=10.5, bold=True, color=accent)
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.1
    set_run(paragraph.add_run(body), size=9.5, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_summary_table(doc):
    counts = Counter(item["status"] for item in ALL_CASES)
    rows = [
        (STATUS_AUTO, counts[STATUS_AUTO], "Existe prueba automática ejecutada y aprobada el 10/08/2026."),
        (STATUS_EVIDENCE, counts[STATUS_EVIDENCE], "La función existe; debe repetirse y adjuntar evidencia manual."),
        (STATUS_EXECUTE, counts[STATUS_EXECUTE], "El caso puede ejecutarse, pero no hay resultado formal registrado."),
        (STATUS_IMPLEMENT, counts[STATUS_IMPLEMENT], "La función o parte exigida todavía no permite ejecutar el caso completo."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Estado documental", "Casos", "Interpretación"]
    widths = [2600, 1000, 5760]
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, size=9, bold=True, color=WHITE)
        set_cell_shading(table.rows[0].cells[index], GREEN)
    repeat_header(table.rows[0])
    for status, count, meaning in rows:
        cells = table.add_row().cells
        fill, color = status_fill(status)
        set_cell_text(cells[0], status, size=8.8, bold=True, color=color)
        set_cell_shading(cells[0], fill)
        set_cell_text(cells[1], count, size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], meaning, size=8.8)
    set_table_geometry(table, widths, indent=0)
    for row in table.rows:
        prevent_row_split(row)
    doc.add_paragraph()


def add_definition_table(doc):
    rows = [
        ("Objetivo", "Demostrar que Trazaap satisface sus requisitos funcionales y no funcionales sin perder integridad entre PostgreSQL, la API, el frontend y Hyperledger Fabric."),
        ("Alcance", "Autenticación, RBAC, proveedores, materias primas, recepción, inspección, productos, órdenes, manufactura, liberación, inventarios, trazabilidad, QR, reportes, blockchain y atributos de calidad."),
        ("Entorno", "Aplicación local con frontend Next.js, backend Node.js/Express, PostgreSQL y red Hyperledger Fabric desplegada en Docker."),
        ("Datos", "Usuarios de cada rol, proveedores, materias primas, recetas, recepciones, órdenes, lotes producidos, liberaciones y casos deliberadamente inválidos."),
        ("Criterio de entrada", "Servicios activos, migraciones aplicadas, usuarios y catálogos disponibles, red Fabric operativa y versión del código identificada."),
        ("Criterio de salida", "Todos los casos aplicables ejecutados; fallos críticos cerrados; evidencias adjuntas; brechas pendientes aceptadas y documentadas."),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, size=9, bold=True, color=DARK_GREEN)
        set_cell_shading(cells[0], LIGHT_GREEN)
        set_cell_text(cells[1], value, size=9)
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, [2100, 7260], indent=0)
    doc.add_paragraph()


def add_execution_rules(doc):
    rules = [
        ("Preparar", "Registrar versión/commit, fecha, navegador, usuario, datos de entrada y estado de los contenedores."),
        ("Ejecutar", "Seguir el escenario, probar el caso válido y las variaciones inválidas indicadas, sin corregir datos durante la misma evidencia."),
        ("Comparar", "Contrastar el resultado observado con el esperado en frontend, respuesta HTTP, PostgreSQL y Fabric cuando aplique."),
        ("Evidenciar", "Adjuntar captura, respuesta de API, consulta SQL o transacción Fabric. La evidencia debe mostrar lote, fecha y resultado."),
        ("Clasificar", "Marcar Aprobada, Fallida, Bloqueada o No aplica. Toda falla debe enlazar un incidente y su posterior re-prueba."),
        ("Cerrar", "Realizar regresión del flujo afectado y obtener aceptación del responsable funcional o director del proyecto."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_text(table.rows[0].cells[0], "Etapa", size=9, bold=True, color=WHITE)
    set_cell_text(table.rows[0].cells[1], "Acción requerida", size=9, bold=True, color=WHITE)
    for cell in table.rows[0].cells:
        set_cell_shading(cell, GREEN)
    repeat_header(table.rows[0])
    for label, value in rules:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, size=8.8, bold=True, color=DARK_GREEN)
        set_cell_text(cells[1], value, size=8.8)
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, [1800, 7560], indent=0)
    doc.add_paragraph()


def add_case_table(doc, cases, repeat_table_header=True):
    headers = ["ID", "Req.", "Caso y ejecución", "Resultado esperado", "Justificación", "Estado / evidencia"]
    widths = [650, 750, 2700, 3100, 4250, 2470]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, size=8.5, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[index], GREEN)
    if repeat_table_header:
        repeat_header(table.rows[0])
    for item in cases:
        cells = table.add_row().cells
        values = [item["id"], item["req"], item["test"], item["expected"], item["why"], f'{item["status"]}\n{item["evidence"]}']
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, size=8.05 if index in (2, 3, 4, 5) else 8.25, bold=index == 0, align=WD_ALIGN_PARAGRAPH.CENTER if index in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT)
        fill, color = status_fill(item["status"])
        set_cell_shading(cells[5], fill)
        for paragraph in cells[5].paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string(color)
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, widths, indent=0)
    doc.add_paragraph()


def add_case_matrix(doc, title, intro, cases, split_at=None, force_new_page=False):
    if force_new_page:
        doc.add_page_break()
    chunks = [cases]
    if split_at and 0 < split_at < len(cases):
        chunks = [cases[:split_at], cases[split_at:]]
    for index, chunk in enumerate(chunks):
        if index > 0:
            doc.add_page_break()
        heading = title if index == 0 else f"{title} (continuación)"
        add_heading(doc, heading, 1)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        text = intro if index == 0 else "Continuación de la matriz. Se repiten los encabezados para conservar el contexto de lectura y ejecución."
        set_run(paragraph.add_run(text), size=9.5, color=GRAY)
        add_case_table(doc, chunk, repeat_table_header=not bool(split_at))


def add_evidence_section(doc):
    add_heading(doc, "6. Evidencias y cierre del plan", 1)
    add_body(doc, "Una prueba se considera terminada únicamente cuando su resultado puede ser revisado por otra persona. Las conversaciones, recuerdos o demostraciones sin registro sirven como exploración, pero no como evidencia formal de aceptación.")
    evidence_rows = [
        ("Interfaz", "Captura completa del formulario, mensaje de éxito/error, lote y fecha visibles."),
        ("API", "Método, endpoint, código HTTP y cuerpo de respuesta sin credenciales sensibles."),
        ("PostgreSQL", "Consulta que demuestre relaciones, movimientos, responsables y ausencia de duplicados."),
        ("Fabric", "txId, función de chaincode, evento consultado, hash y estado de validación."),
        ("Rendimiento", "Herramienta, número de usuarios, volumen, percentiles y condiciones del equipo."),
        ("Usabilidad", "Participantes, tarea, tiempo, errores, observaciones y aceptación del usuario."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(("Tipo", "Evidencia mínima")):
        set_cell_text(table.rows[0].cells[index], header, size=9, bold=True, color=WHITE)
        set_cell_shading(table.rows[0].cells[index], GREEN)
    repeat_header(table.rows[0])
    for label, value in evidence_rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, size=9, bold=True, color=DARK_GREEN)
        set_cell_text(cells[1], value, size=9)
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, [1900, 7460], indent=0)
    doc.add_paragraph()
    add_heading(doc, "6.1 Registro de ejecución recomendado", 2)
    add_body(doc, "Para cada ID de la matriz se debe registrar: fecha, ejecutor, versión del sistema, datos usados, resultado observado, estado final, enlace a evidencia, incidente relacionado y fecha de re-prueba. Los estados del presente documento describen evidencia disponible, no sustituyen el resultado de una ejecución futura.")
    add_heading(doc, "6.2 Evidencia automatizada disponible", 2)
    add_body(doc, "El 10 de agosto de 2026 se ejecutaron 13 pruebas del backend y 10 del chaincode; las 23 finalizaron correctamente. Los comandos reproducibles son `npm.cmd test` en `backend` y en `fabric/chaincode/traceability`.")
    add_callout(doc, "Conclusión de auditoría", "Trazaap cuenta con pruebas automáticas reales para RBAC, códigos externos, normalización/hash y reglas RF13 del chaincode. Para afirmar que todas las pruebas fueron realizadas todavía se deben adjuntar evidencias de los flujos manuales, ejecutar los atributos no funcionales y completar los módulos marcados como pendientes de implementación.", fill=LIGHT_GOLD, accent=GOLD)
    add_heading(doc, "7. Fuentes revisadas", 1)
    sources = [
        "Requerimientos Funcionales y No Funcionales.pdf, versión disponible en el proyecto.",
        "backend/test/*.test.js: pruebas automatizadas del backend.",
        "fabric/chaincode/traceability/test/traceability.test.js: pruebas automatizadas del chaincode.",
        "Código actual de frontend, backend, migraciones y configuración Fabric de Trazaap.",
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, source in enumerate(sources, start=1):
        cells = table.add_row().cells
        set_cell_text(cells[0], f"Fuente {index}", size=9, bold=True, color=DARK_GREEN)
        set_cell_shading(cells[0], LIGHT_GREEN)
        set_cell_text(cells[1], source, size=9)
        prevent_row_split(table.rows[-1])
    set_table_geometry(table, [1600, 7760], indent=0)


def build_document():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    configure_header_footer(section, different_first=True)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(34)
    cover.paragraph_format.space_after = Pt(18)
    if LOGO.exists():
        picture = cover.add_run().add_picture(str(LOGO), width=Inches(1.12))
        picture._inline.docPr.set("descr", "Logo de Trazaap")
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(8)
    set_run(kicker.add_run("TRAZAAP | ASEGURAMIENTO DE CALIDAD"), size=10, bold=True, color=GREEN)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    set_run(title.add_run("Plan integral de pruebas"), size=27, bold=True, color=DARK_GREEN, font="Aptos Display")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    set_run(subtitle.add_run("Validación funcional, blockchain y no funcional del sistema de trazabilidad alimentaria"), size=13, color=GRAY)
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Documento", "Plan de pruebas y matriz de casos"),
        ("Versión", "1.0"),
        ("Fecha de corte", date(2026, 8, 10).strftime("%d/%m/%Y")),
        ("Estado", "Base para ejecución, evidencia y aceptación del equipo"),
    ]
    for index, (label, value) in enumerate(meta_rows):
        set_cell_text(meta.rows[index].cells[0], label, size=9, bold=True, color=DARK_GREEN)
        set_cell_shading(meta.rows[index].cells[0], LIGHT_GREEN)
        set_cell_text(meta.rows[index].cells[1], value, size=9)
        prevent_row_split(meta.rows[index])
    set_table_geometry(meta, [2100, 5100], indent=0)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(25)
    note.paragraph_format.space_after = Pt(0)
    set_run(note.add_run("Documento de trabajo compartible. Los estados reflejan evidencia verificable disponible, no recuerdos de pruebas informales."), size=9, italic=True, color=GRAY)
    doc.add_page_break()

    add_heading(doc, "1. Resumen ejecutivo", 1)
    add_body(doc, f"Este plan define {len(ALL_CASES)} casos necesarios para evaluar Trazaap de extremo a extremo. La matriz cubre requisitos funcionales, reglas de Hyperledger Fabric, seguridad, rendimiento, usabilidad, disponibilidad, cumplimiento e interoperabilidad.")
    add_callout(doc, "Estado comprobable al 10/08/2026", "Se ejecutaron 23 pruebas automáticas: 13 del backend y 10 del chaincode, todas aprobadas. La compilación del frontend también fue exitosa, pero una compilación no reemplaza las pruebas funcionales de interfaz.")
    add_summary_table(doc)
    add_body(doc, "El estado `Pendiente de evidencia` no significa necesariamente que la función nunca se haya usado. Indica que debe repetirse de forma controlada y adjuntarse una evidencia revisable. Esto permite aprovechar las pruebas manuales ya realizadas sin afirmar un cumplimiento que no pueda demostrarse.")

    add_heading(doc, "2. Alcance y criterios", 1)
    add_definition_table(doc)
    add_heading(doc, "2.1 Niveles de prueba", 2)
    add_body(doc, "Las pruebas unitarias comprueban reglas aisladas; las de integración verifican API, PostgreSQL y Fabric; las pruebas de sistema recorren la interfaz completa; las no funcionales miden atributos de calidad; y la aceptación confirma que el flujo representa la operación real de Angela's Bagels.")
    add_heading(doc, "2.2 Regla para declarar una prueba aprobada", 2)
    add_body(doc, "El resultado observado debe coincidir con el esperado, no debe dejar efectos laterales incorrectos y debe existir una evidencia vinculada al ID del caso. Cuando una prueba falla, su corrección exige una re-prueba y una regresión de los módulos relacionados.")

    add_heading(doc, "3. Procedimiento de ejecución", 1)
    add_execution_rules(doc)
    add_callout(doc, "Importante", "Antes de ejecutar casos blockchain se debe comprobar que CA, orderer, ambos peers, canal y chaincode estén activos. Una respuesta pendiente o no encontrada puede deberse a infraestructura apagada y no necesariamente a una falla del formulario.", fill=LIGHT_BLUE, accent=BLUE)

    landscape = doc.add_section(WD_SECTION_START.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width = Inches(11)
    landscape.page_height = Inches(8.5)
    landscape.top_margin = Inches(0.55)
    landscape.bottom_margin = Inches(0.55)
    landscape.left_margin = Inches(0.6)
    landscape.right_margin = Inches(0.6)
    configure_header_footer(landscape, "MATRIZ DE PRUEBAS | TRAZAAP")
    add_case_matrix(doc, "4. Matriz de pruebas funcionales", "Los casos PF recorren la operación completa desde el acceso hasta reportes y módulos normativos pendientes. Cada fila agrupa las variaciones válidas y negativas esenciales del escenario.", functional_cases)
    add_case_matrix(doc, "4.1 Matriz de blockchain e integración", "Los casos BC corresponden a pruebas automáticas ejecutadas sobre normalización, consistencia transaccional y chaincode. Todas se aprobaron en la fecha de corte.", blockchain_cases, split_at=9, force_new_page=True)
    add_case_matrix(doc, "5. Matriz de pruebas no funcionales", "Estas pruebas requieren un entorno estable, herramientas de medición y evidencia cuantitativa. No deben marcarse como cumplidas a partir de una demostración visual.", nonfunctional_cases, split_at=7, force_new_page=True)
    add_callout(doc, "Condiciones para ejecutar pruebas no funcionales", "Registrar hardware, versiones, volumen inicial, número de usuarios concurrentes, duración, herramienta y percentiles. Las pruebas de disponibilidad y consenso requieren provocar fallas controladas y documentar la recuperación.", fill=LIGHT_BLUE, accent=BLUE)

    portrait = doc.add_section(WD_SECTION_START.NEW_PAGE)
    portrait.orientation = WD_ORIENT.PORTRAIT
    portrait.page_width = Inches(8.5)
    portrait.page_height = Inches(11)
    portrait.top_margin = Inches(0.72)
    portrait.bottom_margin = Inches(0.72)
    portrait.left_margin = Inches(0.82)
    portrait.right_margin = Inches(0.82)
    configure_header_footer(portrait)
    add_evidence_section(doc)

    doc.core_properties.title = "Plan integral de pruebas Trazaap"
    doc.core_properties.subject = "Matriz funcional, blockchain y no funcional"
    doc.core_properties.author = "Equipo de proyecto Trazaap"
    doc.core_properties.keywords = "Trazaap, pruebas, trazabilidad, Hyperledger Fabric, calidad"
    doc.save(OUT)
    print(OUT)
    print(f"Casos: {len(ALL_CASES)}")
    print(dict(Counter(item["status"] for item in ALL_CASES)))


if __name__ == "__main__":
    build_document()
