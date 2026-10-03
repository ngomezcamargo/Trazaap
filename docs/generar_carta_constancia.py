from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUTPUT = Path(
    r"C:\Users\ngome\Desktop\Nicolas\Proyecto de Grado\Carta_constancia_revision_tecnica_Trazaap_INVIMA.docx"
)


def configurar_documento(documento: Document) -> None:
    seccion = documento.sections[0]
    seccion.top_margin = Cm(2.5)
    seccion.bottom_margin = Cm(2.5)
    seccion.left_margin = Cm(3)
    seccion.right_margin = Cm(3)

    estilo = documento.styles["Normal"]
    estilo.font.name = "Arial"
    estilo._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    estilo.font.size = Pt(11)


def agregar_parrafo(documento: Document, texto: str = "", despues: int = 10):
    parrafo = documento.add_paragraph()
    parrafo.paragraph_format.space_after = Pt(despues)
    parrafo.paragraph_format.line_spacing = 1.15
    parrafo.add_run(texto)
    return parrafo


def main() -> None:
    documento = Document()
    configurar_documento(documento)

    fecha = agregar_parrafo(documento, "[Ciudad], ____ de ______________ de 2026", despues=24)
    fecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    agregar_parrafo(documento, "A quien corresponda:", despues=18)

    agregar_parrafo(
        documento,
        "Yo, [NOMBRE COMPLETO DE LA INGENIERA], identificada con [TIPO Y NUMERO DE DOCUMENTO], "
        "en mi calidad de Ingeniera de Alimentos, certifico que revisé el software Trazaap, "
        "desarrollado para la gestión de la trazabilidad alimentaria de Angela's Bagels S.A.S.",
        despues=12,
    )

    agregar_parrafo(
        documento,
        "De acuerdo con la revisión realizada, certifico que el software permite registrar, "
        "conservar y consultar la información necesaria para realizar la trazabilidad de las "
        "materias primas, los lotes, la producción, la liberación y el despacho de los productos. "
        "En lo relacionado con el registro y consulta de la trazabilidad, el software cumple con "
        "los requisitos establecidos en la Resolución 2674 de 2013 del Ministerio de Salud y "
        "Protección Social, cuya vigilancia sanitaria corresponde al INVIMA.",
        despues=12,
    )

    agregar_parrafo(
        documento,
        "La presente constancia se expide a solicitud de los interesados.",
        despues=30,
    )

    agregar_parrafo(documento, "Atentamente,", despues=32)
    agregar_parrafo(documento, "_______________________________", despues=4)
    agregar_parrafo(documento, "[NOMBRE COMPLETO DE LA INGENIERA]", despues=4)
    agregar_parrafo(documento, "Ingeniera de Alimentos", despues=4)
    agregar_parrafo(documento, "C.C. [NUMERO]", despues=4)
    agregar_parrafo(documento, "Matrícula profesional [NUMERO]", despues=4)

    documento.core_properties.title = "Constancia de cumplimiento de trazabilidad"
    documento.core_properties.subject = "Revisión del software Trazaap"
    documento.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
