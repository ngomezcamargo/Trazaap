from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Informe_Implementacion_RF13_Trazaap.docx"
LOGO = ROOT / "frontend" / "public" / "trazaap-logo.jpeg"

GREEN_DARK = "075E2D"
GREEN = "16823B"
GREEN_MID = "4B9D52"
GREEN_LIGHT = "EAF5ED"
GREEN_PALE = "F5FAF6"
INK = "17231B"
MUTED = "5C6A61"
GRAY_LIGHT = "F2F4F3"
GRAY_BORDER = "CCD8CF"
WHITE = "FFFFFF"
AMBER = "A86B00"
AMBER_LIGHT = "FFF3D5"
RED = "A52828"
RED_LIGHT = "FDECEC"

USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRAY_BORDER, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    if sum(widths_dxa) != USABLE_DXA:
        raise ValueError(f"Los anchos de tabla deben sumar {USABLE_DXA} DXA")

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, bold=False, color=INK, font="Calibri", italic=False):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)


def set_paragraph_keep(paragraph, keep_with_next=False, keep_together=False):
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = keep_together


def create_numbering_instance(doc, ordered):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if ordered else "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if ordered else "•")
    level.append(lvl_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)

    abstract.append(level)
    first_num_index = next(
        (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list(doc, items, ordered=False):
    num_id = create_numbering_instance(doc, ordered)
    paragraphs = []
    for item in items:
        paragraph = doc.add_paragraph(style="Normal")
        paragraph.paragraph_format.space_after = Pt(4 if not ordered else 5)
        paragraph.paragraph_format.line_spacing = 1.10
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        level = OxmlElement("w:ilvl")
        level.set(qn("w:val"), "0")
        number = OxmlElement("w:numId")
        number.set(qn("w:val"), str(num_id))
        num_pr.append(level)
        num_pr.append(number)
        p_pr.append(num_pr)
        run = paragraph.add_run(item)
        set_run_font(run)
        paragraphs.append(paragraph)
    return paragraphs


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=9, color=MUTED)


def add_num_pages_field(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " NUMPAGES "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)
    set_run_font(run, size=9, color=MUTED)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    set_paragraph_keep(paragraph, keep_with_next=True)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph(style="Normal")
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_run_font(run)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.10
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    set_run_font(run)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.10
    return paragraph


def keep_table_together(table):
    paragraphs = [paragraph for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]
    for paragraph in paragraphs[:-1]:
        paragraph.paragraph_format.keep_with_next = True


def add_code_line(doc, text):
    paragraph = doc.add_paragraph(style="Code")
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=INK, font="Consolas")
    return paragraph


def add_callout(doc, label, text, fill=GREEN_LIGHT, accent=GREEN_DARK):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [USABLE_DXA])
    set_table_borders(table, color=accent, size=8)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, bold=True, color=accent)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths_dxa, header_fill=GREEN_DARK, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, header_fill)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_run_font(run, size=9.3, bold=True, color=WHITE)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        if row_index % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, GREEN_PALE)
        for index, value in enumerate(values):
            paragraph = row.cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)

    trailing = doc.add_paragraph()
    trailing.paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, GREEN_DARK, 16, 8),
        "Heading 2": (13, GREEN, 12, 6),
        "Heading 3": (12, GREEN_DARK, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.font.color.rgb = rgb(INK)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.10

    code = styles["Code"] if "Code" in styles else styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(9)
    code.font.color.rgb = rgb(INK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(2)
    code.paragraph_format.line_spacing = 1.0


def configure_page(doc):
    doc.settings.odd_and_even_pages_header_footer = True
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.82)
        section.bottom_margin = Inches(0.76)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def configure_header_footer(section):
    section.different_first_page_header_footer = True

    def write_header(header):
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        left = paragraph.add_run("TRAZAAP  |  INFORME TÉCNICO RF13")
        set_run_font(left, size=8.5, bold=True, color=MUTED)

    def write_footer(footer):
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_after = Pt(0)
        prefix = paragraph.add_run("Proyecto de grado  |  Página ")
        set_run_font(prefix, size=9, color=MUTED)
        add_page_field(paragraph)
        middle = paragraph.add_run(" de ")
        set_run_font(middle, size=9, color=MUTED)
        add_num_pages_field(paragraph)

    write_header(section.header)
    write_header(section.even_page_header)
    write_footer(section.footer)
    write_footer(section.even_page_footer)

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(10)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    picture = run.add_picture(str(LOGO), width=Inches(1.65))
    picture._inline.docPr.set("title", "Logo de Trazaap")
    picture._inline.docPr.set("descr", "Logo de Trazaap: hojas verdes y bloques de trazabilidad")
    paragraph.paragraph_format.space_after = Pt(10)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(7)
    run = kicker.add_run("REQUERIMIENTO FUNCIONAL RF13")
    set_run_font(run, size=11, bold=True, color=GREEN)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_together = True
    run = title.add_run("Contratos inteligentes\nmediante chaincode")
    set_run_font(run, size=27, bold=True, color=GREEN_DARK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    run = subtitle.add_run("Informe de implementación, integración y validación")
    set_run_font(run, size=14, color=MUTED)

    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120], indent_dxa=0)
    set_table_borders(table, color=GREEN_DARK, size=5)
    metadata = [
        ("ESTADO", "Completado"),
        ("CHAINCODE", "v2.2 / secuencia 4"),
        ("FECHA", "9 de agosto de 2026"),
    ]
    for index, (label, value) in enumerate(metadata):
        cell = table.cell(0, index)
        set_cell_shading(cell, GREEN_LIGHT)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(label)
        set_run_font(run, size=8.5, bold=True, color=GREEN)
        paragraph.add_run("\n")
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=10.5, bold=True, color=GREEN_DARK)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(18)

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.space_after = Pt(4)
    statement.paragraph_format.left_indent = Inches(0.55)
    statement.paragraph_format.right_indent = Inches(0.55)
    run = statement.add_run(
        "PostgreSQL conserva la operación del sistema y Hyperledger Fabric protege la evidencia "
        "criptográfica y las decisiones automáticas de trazabilidad."
    )
    set_run_font(run, size=12, italic=True, color=INK)

    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.paragraph_format.space_before = Pt(20)
    run = brand.add_run("Trazaap | Trazabilidad alimentaria segura")
    set_run_font(run, size=10, bold=True, color=GREEN_DARK)

    doc.add_page_break()


def add_document_content(doc):
    add_heading(doc, "Resumen ejecutivo", 1)
    add_body(
        doc,
        "Se implementó el requerimiento RF13 para que Hyperledger Fabric no se limite a almacenar "
        "hashes, sino que ejecute reglas de negocio mediante el chaincode traceability. La solución "
        "protege la inmutabilidad de los eventos, bloquea despachos no conformes, registra decisiones "
        "auditables, permite confirmar la recepción del cliente y genera alertas de vencimiento."
    )
    add_body(
        doc,
        "La integración conserva la arquitectura híbrida del proyecto: PostgreSQL sigue siendo la "
        "base operativa, mientras que Fabric actúa como capa de integridad, auditoría y confianza. "
        "No se creó una tabla bloque_blockchain ni una simulación del ledger en la base relacional."
    )
    add_callout(
        doc,
        "Resultado principal",
        "Una liberación aprobada solo se convierte en despacho cuando el chaincode autoriza todas las "
        "reglas. Si Fabric no está disponible, el sistema responde HTTP 503 y no modifica el inventario."
    )

    add_heading(doc, "Contenido del informe", 2)
    contents = [
        "Arquitectura y decisiones de diseño.",
        "Funciones implementadas en el chaincode.",
        "Integración con backend, frontend y trazabilidad.",
        "Flujos automáticos de RF13.",
        "Pruebas, despliegue y evidencia de cumplimiento.",
        "Guía de demostración y recomendaciones siguientes.",
    ]
    add_list(doc, contents)

    add_heading(doc, "1. Arquitectura implementada", 1)
    add_body(
        doc,
        "La solución separa claramente los datos operativos de la evidencia inmutable. Esta decisión "
        "evita usar Fabric como base de consultas administrativas y evita que PostgreSQL intente "
        "comportarse como una blockchain."
    )

    architecture_rows = [
        ("PostgreSQL", "Usuarios, proveedores, materias primas, órdenes, manufactura, liberación e inventarios."),
        ("Backend Node.js/Express", "Normaliza payloads, aplica reglas operativas y coordina PostgreSQL con Fabric."),
        ("Chaincode traceability", "Valida reglas, calcula hashes, registra evidencia y decisiones inmutables."),
        ("Hyperledger Fabric", "Ledger permisionado con dos peers, orderer, CA, canal e identidad MSP."),
        ("Frontend Next.js", "Muestra bloqueos, estados blockchain, alertas, confirmaciones y reportes."),
    ]
    add_table(doc, ["Componente", "Responsabilidad"], architecture_rows, [2300, 7060])

    add_heading(doc, "1.1 Flujo de comunicación", 2)
    add_list(doc, [
        "El usuario registra una operación desde el frontend.",
        "El backend valida los datos y guarda la información operativa en PostgreSQL.",
        "El backend construye un payload estable con los datos funcionales del evento.",
        "El chaincode canoniza el payload, calcula SHA-256 y registra la evidencia en Fabric.",
        "La trazabilidad reconstruye el payload actual y solicita al chaincode verificar su integridad.",
        "El frontend muestra VERIFICADO, VERIFICADO_CORREGIDO, ALTERADO o NO_ENCONTRADO.",
    ], ordered=True)

    add_callout(
        doc,
        "Regla de separación",
        "PostgreSQL = operación completa. Hyperledger Fabric = evidencia criptográfica, decisiones y auditoría."
    )

    add_heading(doc, "2. Chaincode implementado", 1)
    add_body(
        doc,
        "El contrato se encuentra en fabric/chaincode/traceability/index.js. La versión desplegada es "
        "2.2 y la definición comprometida en el canal utiliza la secuencia 4."
    )

    chaincode_rows = [
        ("registrarEvento", "Registra un evento nuevo y rechaza cualquier intento de sobrescritura."),
        ("validarEvento", "Compara el payload actual con el original o con la última corrección autorizada."),
        ("registrarCorreccionEvento", "Crea una evidencia nueva vinculada al evento original sin sustituirlo."),
        ("consultarHistorialEvento", "Retorna el original y su secuencia de correcciones inmutables."),
        ("validarDespacho", "Evalúa vencimiento, inventario, manufactura, liberación, transporte y controles críticos."),
        ("registrarDespacho", "Registra el despacho únicamente cuando todas las reglas son conformes."),
        ("confirmarRecepcionCliente", "Registra una recepción única después de validar lote y credenciales externas."),
        ("registrarAlertaVencimiento", "Registra una sola alerta por lote vencido con unidades disponibles."),
        ("consultarEvento / consultarEventosPorLote", "Permite recuperar evidencia individual o reconstruir el historial del lote."),
    ]
    add_table(doc, ["Función", "Qué hace"], chaincode_rows, [3000, 6360], font_size=9.1)

    add_heading(doc, "2.1 Inmutabilidad y correcciones", 2)
    add_body(
        doc,
        "La clave principal de un evento mantiene la forma tipoEvento:idEntidad. Si esa clave ya existe, "
        "registrarEvento retorna EVENTO_DUPLICADO. Por tanto, ningún flujo normal puede reemplazar el "
        "hash ni el payload originalmente comprometido."
    )
    add_body(
        doc,
        "Cuando una modificación es legítima, el gerente registra una corrección con motivo, actor, "
        "payload corregido, hash, txId, timestamp e identidad Fabric. validarEvento compara el estado "
        "operativo con la corrección más reciente y retorna VERIFICADO_CORREGIDO si coincide."
    )
    add_callout(
        doc,
        "Diferencia importante",
        "Una modificación directa en PostgreSQL aparece como ALTERADO. Solo una corrección explícita y "
        "auditada puede establecer una nueva versión válida."
    )

    add_heading(doc, "2.2 Identidad y metadatos Fabric", 2)
    add_body(
        doc,
        "El chaincode no confía únicamente en el nombre de actor enviado por el backend. Cada registro "
        "incluye también txId, timestamp de la transacción, MSP del invocador e identidad Fabric obtenida "
        "desde el contexto de la red."
    )

    add_heading(doc, "3. Integración del backend", 1)
    add_body(
        doc,
        "La comunicación con Fabric se mantuvo desacoplada en el módulo backend/src/modulos/blockchain. "
        "fabric.client.js traduce los errores del Gateway y fabric-traceability.service.js ejecuta las "
        "transacciones y consultas del contrato. blockchain.service.js construye eventos, valida su estado "
        "y administra las correcciones."
    )

    add_heading(doc, "3.1 Payloads estables", 2)
    add_body(
        doc,
        "Cada tipo de evento tiene un constructor específico dentro de backend/src/modulos/blockchain/payloads. "
        "Estos constructores consultan las relaciones necesarias, excluyen metadatos técnicos y ordenan las "
        "claves para que la misma información produzca siempre el mismo hash."
    )
    payload_rows = [
        ("Sí entran", "Lotes, productos, proveedores, responsables, cantidades, tiempos, temperaturas, decisiones y observaciones funcionales."),
        ("No entran", "created_at, updated_at, tokens, estados de interfaz, logs e identificadores técnicos dentro del payload."),
        ("Se usa aparte", "idEntidad se utiliza para localizar el evento en Fabric, sin contaminar el contenido funcional."),
    ]
    add_table(doc, ["Categoría", "Tratamiento"], payload_rows, [2100, 7260])

    add_heading(doc, "3.2 Errores controlados", 2)
    add_list(doc, [
        "Fabric no disponible: HTTP 503 en operaciones críticas.",
        "Evento duplicado: se informa sin sobrescribir la evidencia.",
        "Lote no encontrado: error funcional controlado.",
        "Despacho bloqueado: retorna todos los motivos y reglas evaluadas.",
        "Recepción ya confirmada y alerta duplicada: se rechazan de forma explícita.",
    ])

    add_heading(doc, "4. Validación y registro del despacho", 1)
    add_body(
        doc,
        "En Trazaap, la liberación aprobada representa actualmente el despacho. Antes de escribir una "
        "liberación aprobada, el backend construye los datos de validación y llama a validarDespacho. "
        "El chaincode recalcula cada regla y emite APROBADO o BLOQUEADO."
    )

    dispatch_rules = [
        ("Manufactura", "Existe manufactura y el lote coincide con el producto fabricado."),
        ("Estado", "La liberación está aprobada y el lote no fue despachado antes."),
        ("Vencimiento", "La fecha de vencimiento es posterior al momento de despacho."),
        ("Liberación", "Etiqueta, lote, fecha, empaque, envase y estado del producto cumplen."),
        ("Transporte", "Conductor, placa, limpieza del vehículo y documentación cumplen."),
        ("Producción", "Fermentación, horneado e inmersión están dentro de los rangos calculados."),
        ("Inventario", "Las unidades disponibles son suficientes para la cantidad despachada."),
    ]
    add_table(doc, ["Grupo de reglas", "Validación realizada"], dispatch_rules, [2400, 6960])

    add_heading(doc, "4.1 Transacción operativa", 2)
    add_list(doc, [
        "El backend solicita la decisión del chaincode antes de escribir en PostgreSQL.",
        "Si la decisión es BLOQUEADO, conserva el formulario y muestra todos los motivos.",
        "Si la decisión es APROBADO, inicia una transacción PostgreSQL.",
        "Registra la liberación y el despacho en Fabric.",
        "Crea o actualiza el inventario de producto terminado.",
        "Confirma la transacción únicamente si Fabric respondió correctamente.",
    ], ordered=True)
    add_body(
        doc,
        "Si registrarDespacho falla, la transacción relacional se revierte. Esto evita liberaciones "
        "parciales, inventarios descontados sin evidencia o lotes marcados como entregados sin autorización."
    )

    add_heading(doc, "4.2 Controles críticos", 2)
    add_body(
        doc,
        "Las tolerancias se calculan en controles-criticos.service.js a partir de los valores estándar del "
        "producto y los valores reales de manufactura. El frontend no define reglas propias. El chaincode "
        "vuelve a verificar numéricamente cada mínimo, máximo y valor recibido."
    )
    control_rows = [
        ("Fermentación", "Tiempo: ±10 %, mínimo 3 min. Temperatura: ±3 °C."),
        ("Horneado", "Tiempo: ±10 %, mínimo 2 min. Temperatura: ±8 °C."),
        ("Inmersión", "Solo para Bagel. Tiempo: ±15 %, mínimo 1 min. Temperatura: ±5 °C."),
    ]
    add_table(doc, ["Proceso", "Tolerancia inicial"], control_rows, [2500, 6860])

    add_heading(doc, "5. Confirmación del cliente", 1)
    add_body(
        doc,
        "El cliente receptor no necesita una cuenta interna. Desde el portal público del QR puede consultar "
        "el lote y confirmar la recepción mediante el número de factura o el código de cliente."
    )
    add_list(doc, [
        "El backend valida las credenciales contra la información operativa del lote.",
        "El chaincode comprueba que el lote existe y posee un despacho aprobado.",
        "Se valida que la factura o el código correspondan al despacho.",
        "Se rechaza cualquier confirmación duplicada.",
        "Fabric registra RECIBIDO_POR_CLIENTE con fecha, actor y txId.",
    ], ordered=True)
    add_callout(
        doc,
        "Privacidad",
        "El portal del consumidor no expone hashes, identidades Fabric ni información comercial sensible. "
        "El detalle técnico permanece reservado para usuarios autorizados y el acceso de auditoría."
    )

    add_heading(doc, "6. Alertas de vencimiento", 1)
    add_body(
        doc,
        "Un chaincode no puede despertarse por sí mismo cuando pasa el tiempo. Por eso se agregó una tarea "
        "programada en vencimientos.job.js. El backend revisa periódicamente los lotes aprobados que ya "
        "vencieron, conservan unidades y no tienen despacho confirmado."
    )
    add_body(
        doc,
        "Cuando encuentra un caso, invoca registrarAlertaVencimiento. El chaincode evita duplicados y "
        "registra lote, producto, fecha de vencimiento, unidades disponibles, fecha de detección y estado "
        "VENCIDO_SIN_DESPACHO. El dashboard gerencial presenta las alertas activas."
    )

    add_heading(doc, "7. Cambios en frontend, trazabilidad y reportes", 1)
    frontend_rows = [
        ("Liberación", "Muestra los motivos del chaincode dentro del formulario y conserva los datos para corregirlos."),
        ("Portal QR", "Permite consulta pública y confirmación controlada por factura o código."),
        ("Dashboard", "Presenta alertas de lotes vencidos sin despacho."),
        ("Trazabilidad interna", "Muestra estado, hashes, mensaje, correcciones y decisiones automáticas."),
        ("Reporte PDF", "Incluye validación, decisión, motivos, fecha Fabric y txId por etapa."),
    ]
    frontend_table = add_table(doc, ["Vista", "Comportamiento agregado"], frontend_rows, [2200, 7160])
    keep_table_together(frontend_table)
    add_body(
        doc,
        "No se creó una pantalla separada para usar blockchain. La integración ocurre automáticamente "
        "dentro de los flujos operativos, tal como exige la arquitectura del proyecto."
    )

    add_heading(doc, "8. Eventos cubiertos", 1)
    event_rows = [
        ("recepcion_materia_prima", "Registro inicial e integridad de la recepción."),
        ("inspeccion_recepcion", "Resultado, condiciones y decisión de aceptación."),
        ("orden_produccion", "Creación y versiones funcionales de la orden."),
        ("producto_fabricado_configurado", "Ficha técnica, variante y receta vigente."),
        ("registro_manufactura", "Lote producido, cantidades, tiempos y temperaturas reales."),
        ("liberacion_producto", "Validaciones y decisión de liberación."),
        ("despacho_producto", "Decisión automática y salida aprobada."),
        ("inventario_producto_terminado", "Disponibilidad posterior a liberación."),
        ("inventario_materia_prima", "Versiones funcionales del saldo de materias primas."),
        ("movimiento_inventario", "Entradas, salidas y referencias operativas."),
        ("confirmacion_recepcion_cliente", "Entrega confirmada por el receptor."),
        ("correccion_evento", "Corrección vinculada al evento original."),
        ("alerta_vencimiento", "Producto vencido con unidades y sin despacho."),
    ]
    add_table(doc, ["Evento Fabric", "Propósito"], event_rows, [3400, 5960], font_size=8.9)

    add_heading(doc, "9. Pruebas realizadas", 1)
    add_body(
        doc,
        "La implementación se verificó con pruebas automatizadas, consultas reales a la red, compilación "
        "del frontend y revisión visual de las vistas."
    )
    test_rows = [
        ("Chaincode", "10 de 10", "Inmutabilidad, alteración, correcciones, despacho, recepción y alertas."),
        ("Backend", "13 de 13", "HTTP 503, rollback, inventario, credenciales públicas, hashes y RBAC."),
        ("Frontend", "17 rutas", "Compilación optimizada sin errores de tipos ni renderizado."),
        ("Red Fabric", "Correcto", "Chaincode 2.2 / secuencia 4 confirmado desde peer0 y peer1."),
        ("Prueba visual", "Correcto", "Portal público, vista técnica, historial de correcciones y reporte."),
    ]
    add_table(doc, ["Capa", "Resultado", "Cobertura"], test_rows, [1800, 1500, 6060], font_size=9)

    add_heading(doc, "9.1 Casos de chaincode", 2)
    add_list(doc, [
        "Registrar un evento nuevo con identidad Fabric.",
        "Rechazar la sobrescritura de un evento existente.",
        "Detectar un registro operativo alterado.",
        "Registrar una corrección conservando el original.",
        "Confirmar una recepción y evitar la confirmación duplicada.",
        "Permitir un despacho conforme y bloquear controles fuera de rango.",
        "Bloquear un lote vencido y evitar alertas repetidas.",
    ])

    add_heading(doc, "10. Despliegue de la red", 1)
    network_rows = [
        ("Organización", "Org1MSP"),
        ("CA", "ca.trazaap.local : 7054"),
        ("Orderer", "orderer.trazaap.local : 7050"),
        ("Peer principal", "peer0.org1.trazaap.local : 7051"),
        ("Peer secundario", "peer1.org1.trazaap.local : 8051"),
        ("Canal", "trazabilidad-channel"),
        ("Chaincode", "traceability 2.2, secuencia 4"),
    ]
    add_table(doc, ["Elemento", "Configuración activa"], network_rows, [3000, 6360])
    add_body(
        doc,
        "El paquete se instaló en ambos peers, fue aprobado para Org1MSP y comprometido en el canal. "
        "El despliegue no eliminó el ledger, no regeneró certificados y no recreó el canal."
    )

    add_heading(doc, "11. Matriz de aceptación de RF13", 1)
    acceptance_rows = [
        ("1", "Un evento existente no puede sobrescribirse.", "Completado", "EVENTO_DUPLICADO y prueba automatizada."),
        ("2", "Una modificación aparece como ALTERADO.", "Completado", "Prueba unitaria y caso real de recepción 42."),
        ("3", "Una corrección conserva el original.", "Completado", "Historial y VERIFICADO_CORREGIDO."),
        ("4", "Un despacho fuera de rango se bloquea.", "Completado", "Reglas y motivos retornados por chaincode."),
        ("5", "Un despacho conforme se registra en Fabric.", "Completado", "Prueba de registrarDespacho."),
        ("6", "Inventario cambia solo tras aprobación.", "Completado", "Transacción y pruebas de no escritura."),
        ("7", "Cliente confirma sin cuenta interna.", "Completado", "Portal público con factura o código."),
        ("8", "Lote vencido genera una sola alerta.", "Completado", "Job periódico y deduplicación Fabric."),
        ("9", "Trazabilidad y reporte muestran decisiones.", "Completado", "Vista técnica y reporte verificados."),
        ("10", "Pruebas automatizadas pasan.", "Completado", "10/10 chaincode y 13/13 backend."),
        ("11", "Red funcional con ambos peers.", "Completado", "Versión 2.2 / secuencia 4 en peer0 y peer1."),
        ("12", "Frontend y backend continúan funcionando.", "Completado", "Build, health check y navegación real."),
    ]
    add_table(doc, ["N.º", "Criterio", "Estado", "Evidencia"], acceptance_rows, [600, 3540, 1500, 3720], font_size=8.4)

    add_heading(doc, "12. Evidencia real de la validación", 1)
    add_body(
        doc,
        "Durante la verificación final se registró una corrección autorizada de la ficha técnica del "
        "producto con id 1. Fabric conservó el evento original, registró una nueva transacción y la "
        "trazabilidad pasó a VERIFICADO_CORREGIDO."
    )
    evidence_rows = [
        ("Entidad", "producto_fabricado_configurado:1"),
        ("Estado", "VERIFICADO_CORREGIDO"),
        ("Motivo", "Actualización autorizada de la ficha técnica y receta vigente."),
        ("Transacción", "8258a46116a474514fd27fa8d13d126040064591e1f1c6759112b972f01ba4ef"),
        ("Original", "Conservado sin modificación."),
    ]
    add_table(doc, ["Dato", "Resultado"], evidence_rows, [2200, 7160], font_size=8.9)
    add_callout(
        doc,
        "Comportamiento correcto",
        "La recepción histórica 42 permanece como ALTERADO porque su información ya no coincide con el "
        "ledger y no existe una corrección justificada. El sistema no oculta esa diferencia."
    )

    add_heading(doc, "13. Cómo demostrar RF13", 1)
    add_list(doc, [
        "Iniciar Docker Desktop y comprobar que CA, orderer, peer0 y peer1 están activos.",
        "Iniciar backend en el puerto 4000 y frontend en el puerto 3000.",
        "Consultar un lote producido desde Trazabilidad.",
        "Revisar la tabla Validación blockchain y el historial técnico de correcciones.",
        "Abrir el reporte PDF y comprobar estados, decisiones, fecha Fabric y txId.",
        "Intentar liberar un lote con un control fuera de rango y verificar el bloqueo.",
        "Corregir los datos, aprobar el despacho y comprobar el registro Fabric.",
        "Abrir el portal QR y confirmar la recepción con factura o código de cliente.",
    ], ordered=True)

    add_heading(doc, "13.1 Comandos de verificación", 2)
    add_code_line(doc, "cd fabric/chaincode/traceability")
    add_code_line(doc, "npm test")
    add_code_line(doc, "cd ../../../backend && npm test")
    add_code_line(doc, "cd ../frontend && npm run build")
    add_body(
        doc,
        "Para actualizar el chaincode sobre un ledger existente se debe incrementar versión y secuencia. "
        "No se debe ejecutar clean.sh ni eliminar certificados para una actualización normal."
    )

    add_heading(doc, "14. Archivos principales", 1)
    file_rows = [
        ("Chaincode", "fabric/chaincode/traceability/index.js"),
        ("Pruebas chaincode", "fabric/chaincode/traceability/test/traceability.test.js"),
        ("Cliente Fabric", "backend/src/modulos/blockchain/fabric.client.js"),
        ("Servicio blockchain", "backend/src/modulos/blockchain/blockchain.service.js"),
        ("Liberación/despacho", "backend/src/modulos/liberacion/liberacion.service.js"),
        ("Controles críticos", "backend/src/modulos/liberacion/controles-criticos.service.js"),
        ("Alertas", "backend/src/modulos/liberacion/vencimientos.job.js"),
        ("Portal público", "backend/src/modulos/publico y frontend/src/app/verificar/[lote]"),
        ("Trazabilidad", "backend/src/modulos/trazabilidad y frontend/src/modulos/trazabilidad"),
        ("Despliegue", "fabric/scripts/deploy-chaincode.sh y fabric/scripts/env.sh"),
    ]
    add_table(doc, ["Responsabilidad", "Ubicación"], file_rows, [2600, 6760], font_size=8.8)

    add_heading(doc, "15. Conclusión", 1)
    add_body(
        doc,
        "RF13 quedó implementado como una capa real de contratos inteligentes sobre Hyperledger Fabric. "
        "El chaincode ya no es un almacenamiento pasivo de hashes: valida reglas, toma decisiones, "
        "protege la inmutabilidad y conserva evidencia auditable de cada resultado."
    )
    add_body(
        doc,
        "La solución mantiene la separación correcta entre la operación empresarial y la confianza "
        "criptográfica. Esto permite que Trazaap continúe usando PostgreSQL para su funcionamiento diario "
        "y utilice Fabric para demostrar integridad, trazabilidad y cumplimiento ante revisiones internas "
        "o procesos de auditoría."
    )
    add_callout(
        doc,
        "Estado final",
        "RF13 completado, probado y desplegado con dos peers, orderer, CA, canal y chaincode funcionales."
    )


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    configure_header_footer(doc.sections[0])

    properties = doc.core_properties
    properties.title = "Informe de implementación RF13 - Trazaap"
    properties.subject = "Contratos inteligentes mediante chaincode"
    properties.author = "Proyecto Trazaap"
    properties.keywords = "Trazaap, Hyperledger Fabric, chaincode, RF13, trazabilidad"
    properties.comments = "Documento técnico de implementación y validación."

    add_cover(doc)
    add_document_content(doc)

    configure_page(doc)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
