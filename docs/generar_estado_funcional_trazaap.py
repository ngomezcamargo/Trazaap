from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parents[1]
OUT = BASE / "docs" / "Estado_funcional_Trazaap.docx"
LOGO = BASE / "frontend" / "public" / "trazaap-logo.jpeg"


GREEN = "0B6B35"
LIGHT_GREEN = "EAF6EE"
MID_GREEN = "2E7D32"
YELLOW = "FFF3CD"
RED = "FDECEC"
BLUE = "EAF2F8"
GRAY = "F5F7F8"
TEXT = "111827"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, size=8.5, color=TEXT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def style_table(table, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    if header:
        for cell in table.rows[0].cells:
            set_cell_shading(cell, GREEN)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(8.5)


def add_table(doc, headers, rows, widths=None, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8.5, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], GREEN)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8)
            if status_col is not None and i == status_col:
                status = str(value).lower()
                if "implementado" in status and "parcial" not in status:
                    set_cell_shading(cells[i], LIGHT_GREEN)
                elif "parcial" in status:
                    set_cell_shading(cells[i], YELLOW)
                elif "pendiente" in status:
                    set_cell_shading(cells[i], RED)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(GREEN)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.size = Pt(9)


def add_code_paths(doc, paths):
    for path, desc in paths:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(path)
        r1.font.name = "Consolas"
        r1.font.size = Pt(8.5)
        r1.bold = True
        r2 = p.add_run(f": {desc}")
        r2.font.size = Pt(8.5)


functional_rows = [
    ("RF01-A", "Implementado", "Recepcion de materias primas", "Se registran proveedor, materia prima, cantidades, lotes, fechas, temperatura, peso, estado y responsable. Tambien se actualiza inventario y se registra evidencia Fabric."),
    ("RF01-B", "Implementado", "Inspeccion y decision de aceptacion", "La inspeccion queda ligada a la recepcion; incluye producto, transporte, vehiculo/conductor como checks funcionales, observaciones y decision final."),
    ("RF02", "Implementado parcial", "Eventos de fabricacion", "Existe registro de manufactura con tiempos y temperaturas reales, lote producido, unidades y responsable. Falta modelar equipos utilizados si se exige literalmente."),
    ("RF03-A", "Implementado parcial", "Envasado y embalado", "La informacion de empaque quedo integrada en liberacion de producto. No existe un modulo separado de envasado/embalado."),
    ("RF03-B", "Implementado parcial", "Identificador unico de lote", "El lote producido se captura, se enlaza a manufactura, liberacion, inventario, trazabilidad y QR. Falta generacion automatica con codigo de fabrica/fechas."),
    ("RF04", "Pendiente", "Almacenamiento", "No hay modulo formal para ubicacion, temperatura de almacenamiento, fechas de ingreso/salida o condiciones por lote."),
    ("RF05", "Implementado parcial", "Distribucion y comercializacion", "La liberacion funciona como despacho operativo: factura, conductor, placa y condiciones del vehiculo. Falta cliente receptor y temperatura de transporte como modulo de despacho independiente."),
    ("RF06", "Implementado parcial", "Trazabilidad bidireccional", "La consulta por lote final reconstruye origen, produccion, liberacion, inventario y validacion blockchain. Falta busqueda completamente hacia adelante desde cualquier materia prima hasta clientes."),
    ("RF07", "Implementado parcial", "Puntos criticos HACCP", "Manufactura compara tiempos/temperaturas estandar vs reales y marca desviaciones. Falta modulo HACCP completo con alertas formales por variable critica."),
    ("RF08", "Implementado parcial", "Calidad e inocuidad", "Recepcion, inspeccion y liberacion documentan decisiones de aceptacion, retencion o rechazo. Faltan controles fisicoquimicos/microbiologicos como formularios independientes."),
    ("RF09", "Implementado", "Bloques/eventos inmutables", "Fabric registra eventos criticos con payload canonico, hash SHA-256, timestamp de transaccion, estado y busqueda por lote."),
    ("RF10", "Implementado parcial", "Codigo QR de lote", "Hay QR en reportes y portal publico por lote final. Falta certificar medidas de etiqueta y completar datos de entrega a cliente."),
    ("RF11", "Implementado parcial", "Portal web publico QR", "Existe /verificar/[lote] sin autenticacion, con resumen publico de trazabilidad y validacion. Falta separar formalmente vista consumidor vs INVIMA."),
    ("RF12", "Implementado parcial", "Control de acceso por roles", "El sistema maneja administrador, gerente y operario con JWT y guardas. No estan los roles cliente, INVIMA y consumidor autenticado."),
    ("RF13", "Implementado parcial", "Contratos inteligentes", "Existe chaincode de Fabric para registrar, validar y consultar eventos. Faltan reglas automaticas: confirmacion cliente, bloqueo por HACCP y alerta por vencimiento."),
    ("RF14", "Implementado parcial", "Reportes de trazabilidad INVIMA", "Hay modulo de reportes y PDF con validacion blockchain/QR. Falta exportacion Excel y generacion masiva hasta 50 lotes."),
    ("RF15", "Pendiente", "Plan de saneamiento", "No hay modulo de limpieza, desinfeccion, plagas o residuos."),
    ("RF16", "Pendiente", "Devoluciones y lotes no conformes", "No hay flujo de devoluciones, retiro, reproceso o destruccion."),
    ("RF17", "Implementado parcial", "Registro de proveedores", "Existe CRUD de proveedores, certificaciones y relacion con materias primas/recepciones. Falta gestion documental completa de certificados/fichas tecnicas."),
]

nonfunctional_rows = [
    ("RNF01", "Implementado parcial", "Seguridad e integridad", "Fabric permite detectar alteraciones mediante hash, pero la red esta en ambiente local de desarrollo."),
    ("RNF02", "Pendiente", "Disponibilidad", "No hay monitoreo, despliegue productivo ni evidencia de 99% mensual."),
    ("RNF03", "Implementado parcial", "Trazabilidad temporal verificable", "Los eventos tienen fechas operativas y timestamp blockchain; falta auditoria formal de secuencia completa."),
    ("RNF04", "Implementado parcial", "Cumplimiento Res. 2674/2013", "Los formularios principales cubren trazabilidad de origen, proceso y liberacion; falta matriz normativa completa dentro del sistema."),
    ("RNF05", "Pendiente", "Rendimiento", "No hay pruebas formales que demuestren <=3 s en consulta o <=10 s en confirmacion blockchain."),
    ("RNF06", "Implementado parcial", "Usabilidad", "Se mejoraron formularios modales, validaciones y cierre automatico; falta prueba controlada de usabilidad."),
    ("RNF07", "Pendiente", "Portabilidad/offline", "La web es responsive, pero no existe modo offline con sincronizacion."),
    ("RNF08", "Implementado parcial", "Confidencialidad diferenciada", "El QR publico limita informacion sensible, pero falta perfil INVIMA y politica de campos por actor."),
    ("RNF09", "Implementado parcial", "Confiabilidad del consenso", "Hay CA, orderer y dos peers en Docker; no hay cluster multi-orderer ni prueba de tolerancia a fallos."),
    ("RNF10", "Pendiente", "Escalabilidad", "No hay pruebas de carga ni arquitectura productiva dimensionada."),
    ("RNF11", "Implementado parcial", "Mantenibilidad", "El codigo esta modularizado por dominios y hay schema consolidado; falta documentacion tecnica completa y versionamiento formal de contratos."),
    ("RNF12", "Implementado parcial", "Interoperabilidad", "La API REST JSON existe con JWT. Falta OAuth 2.0 y compatibilidad GS1 EPCIS 2.0."),
]

module_rows = [
    ("Autenticacion y sesion", "backend/src/modulos/autenticacion, frontend/src/modulos/autenticacion, frontend/src/utilidades/sesion.js", "Login por POST, JWT, perfil, listado de operarios, cierre por inactividad a las 2 horas."),
    ("Roles y permisos", "backend/src/middlewares/roles.middleware.js, frontend/src/comunes/GuardiaRol.js, frontend/src/utilidades/roles.js", "Permite controlar vistas y endpoints para administrador, gerente y operario."),
    ("Proveedores", "backend/src/modulos/proveedores, frontend/src/app/proveedores", "CRUD de proveedores, datos comerciales, NIT, contacto, certificaciones y relacion con materias primas."),
    ("Materias primas", "backend/src/modulos/materias_primas, frontend/src/app/materias-primas", "Gestion del catalogo de insumos, unidades, condiciones y proveedor asociado."),
    ("Recepcion e inspeccion", "backend/src/modulos/recepciones, frontend/src/modulos/recepciones/FormularioRecepcion.js", "Registra recepcion, inspeccion, decision final, inventario inicial y evidencia blockchain."),
    ("Inventario de materias primas", "backend/src/modulos/inventario, backend/src/modulos/recepciones, backend/src/modulos/produccion", "Se alimenta por recepciones aceptadas y se descuenta por manufactura real."),
    ("Productos y recetas", "backend/src/modulos/produccion, frontend/src/modulos/produccion/FormularioOrdenProduccion.js", "Define productos fabricados, variantes por tamano/presentacion, receta e insumos por unidad."),
    ("Ordenes de produccion", "backend/src/modulos/produccion, frontend/src/app/produccion/page.js", "Planifica productos, cantidades y materias primas requeridas; no registra datos reales de manufactura."),
    ("Registro de manufactura", "backend/src/modulos/produccion/produccion.service.js, frontend/src/modulos/produccion/FormularioOrdenProduccion.js", "Registra lo ocurrido: lote producido, unidades, tiempos, temperaturas, responsable y consumos reales."),
    ("Liberacion de producto", "backend/src/modulos/liberacion, frontend/src/app/liberacion, frontend/src/modulos/liberacion/FormularioLiberacion.js", "Control final del producto terminado; si se aprueba crea inventario terminado y queda listo para salida."),
    ("Inventario producto terminado", "backend/src/modulos/inventario, backend/src/modulos/liberacion", "Se genera al aprobar liberacion; relaciona lote, producto, unidades, vencimiento y estado."),
    ("Trazabilidad", "backend/src/modulos/trazabilidad, frontend/src/app/trazabilidad", "Consulta por lote final u origen y agrega recepcion, inspeccion, produccion, manufactura, liberacion y validacion Fabric."),
    ("Reportes", "frontend/src/app/reportes, frontend/src/app/reportes/trazabilidad/[lote]/page.js", "Genera reporte PDF visual con eventos, estados blockchain y codigos QR de verificacion."),
    ("Portal publico QR", "backend/src/modulos/publico, frontend/src/app/verificar/[lote]/page.js", "Permite consultar trazabilidad publica sin autenticacion a partir del lote producido."),
    ("Blockchain Fabric", "backend/src/modulos/blockchain, fabric/chaincode/traceability, fabric/scripts", "Construye payloads funcionales, registra y valida evidencias en Fabric mediante chaincode."),
]


def build_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("Trazaap - Estado funcional del software")
    hr.font.size = Pt(8)
    hr.font.color.rgb = RGBColor.from_string(GREEN)

    title_table = doc.add_table(rows=1, cols=2)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    title_cells = title_table.rows[0].cells
    if LOGO.exists():
        title_cells[0].paragraphs[0].add_run().add_picture(str(LOGO), width=Cm(2.2))
    set_cell_text(title_cells[1], "Estado funcional del software Trazaap\nChecklist de requerimientos, ubicacion en codigo y justificacion tecnica", bold=True, size=16, color=GREEN)
    title_cells[0].width = Cm(2.7)
    title_cells[1].width = Cm(15)
    doc.add_paragraph("Fecha de corte: 18 de julio de 2026. Documento elaborado con base en la propuesta inicial, los requerimientos funcionales/no funcionales y la revision del codigo actual.", style=None)

    add_heading(doc, "1. Checklist de cumplimiento", 1)
    doc.add_paragraph("La siguiente matriz separa lo implementado, lo parcial y lo pendiente. 'Parcial' significa que existe una base funcional en el sistema, pero aun falta una parte del requerimiento literal o una validacion formal para darlo por cerrado.")
    add_heading(doc, "1.1 Requerimientos funcionales", 2)
    add_table(doc, ["ID", "Estado", "Funcion", "Evidencia actual / brecha"], functional_rows, widths=[1.4, 2.3, 3.7, 10.5], status_col=1)

    add_heading(doc, "1.2 Requerimientos no funcionales", 2)
    add_table(doc, ["ID", "Estado", "Atributo", "Evidencia actual / brecha"], nonfunctional_rows, widths=[1.4, 2.3, 3.7, 10.5], status_col=1)

    add_heading(doc, "2. Ubicacion de funcionalidades en el codigo", 1)
    add_table(doc, ["Modulo", "Ubicacion principal", "Que hace"], module_rows, widths=[3.2, 6.8, 8.0])

    add_heading(doc, "3. Explicacion por capas", 1)
    add_heading(doc, "3.1 Frontend web", 2)
    doc.add_paragraph("El frontend esta construido en Next.js y organiza la experiencia por vistas del sistema. La carpeta frontend/src/app contiene las rutas principales: panel, proveedores, materias primas, recepciones, produccion, liberacion, inventario, trazabilidad, reportes y verificacion publica por QR. Las carpetas frontend/src/modulos y frontend/src/comunes contienen formularios reutilizables, guardas de sesion y componentes de layout.")
    add_code_paths(doc, [
        ("frontend/src/app/produccion/page.js", "vista principal de produccion con ordenes, orden activa, manufactura y productos."),
        ("frontend/src/modulos/produccion/FormularioOrdenProduccion.js", "formulario grande de ordenes, productos, recetas, materias y manufactura."),
        ("frontend/src/app/liberacion/page.js", "vista del nuevo segmento Liberacion."),
        ("frontend/src/app/trazabilidad/page.js", "entrada privada para consultar trazabilidad por lote."),
        ("frontend/src/app/verificar/[lote]/page.js", "portal publico que se abre desde el QR."),
        ("frontend/src/app/reportes/trazabilidad/[lote]/page.js", "reporte imprimible/PDF con validacion blockchain."),
    ])

    add_heading(doc, "3.2 Servicios frontend y comunicacion HTTP", 2)
    doc.add_paragraph("Los archivos de frontend/src/servicios encapsulan las llamadas a la API. Esto evita que cada pantalla arme manualmente fetch, headers y manejo de errores. El archivo api.js agrega el token JWT cuando existe, registra actividad de sesion y limpia la sesion si el backend responde 401 o si pasan mas de dos horas de inactividad.")
    add_code_paths(doc, [
        ("frontend/src/servicios/api.js", "cliente HTTP comun para GET, POST, PUT y DELETE."),
        ("frontend/src/utilidades/sesion.js", "guarda token, usuario y ultima actividad; define LIMITE_INACTIVIDAD_MS = 2 horas."),
        ("frontend/src/comunes/GuardiaSesion.js", "protege vistas privadas y programa cierre automatico por inactividad."),
        ("frontend/src/comunes/GuardiaRol.js", "bloquea vistas si el rol no corresponde."),
    ])

    add_heading(doc, "3.3 Backend Express", 2)
    doc.add_paragraph("El backend centraliza reglas de negocio, validacion, autenticacion, acceso a PostgreSQL y comunicacion con Fabric. El punto de entrada es backend/src/app.js: carga middleware de seguridad, CORS, JSON, logs, rutas de dominio y manejo uniforme de errores. Cada modulo sigue una separacion clara: routes recibe la URL, controller traduce request/response, schemas valida entrada, service contiene la regla de negocio y repository ejecuta SQL.")
    add_code_paths(doc, [
        ("backend/src/app.js", "monta todas las rutas bajo el prefijo de API."),
        ("backend/src/server.js", "prueba conexion PostgreSQL y levanta el servidor."),
        ("backend/src/middlewares/autenticarJwt.js", "valida tokens en endpoints protegidos."),
        ("backend/src/middlewares/roles.middleware.js", "controla permisos por rol."),
        ("backend/src/middlewares/validarSolicitud.js", "aplica schemas antes de llegar al servicio."),
    ])

    add_heading(doc, "3.4 Base de datos PostgreSQL", 2)
    doc.add_paragraph("PostgreSQL conserva la informacion operativa completa. La decision es correcta porque el sistema necesita relaciones, consultas, actualizaciones e inventarios que serian costosos e incomodos en blockchain. La blockchain no reemplaza la base de datos; solo conserva evidencia criptografica verificable de eventos criticos.")
    add_bullets(doc, [
        "El schema consolidado actual vive en backend/sql/001_schema_actual.sql y backend/documentacion/schema_actual_consolidado.sql.",
        "Las tablas principales son roles, users, providers, raw_materials, receptions, reception_inspections, inventario_materias_primas, inventario_movimientos, productos_fabricados, producto_variantes, producto_variante_materia_prima, ordenes_produccion, ordenes_produccion_productos, ordenes_produccion_materias, registro_manufactura, liberacion_producto e inventario_producto_terminado.",
        "Las llaves foraneas enlazan proveedor, materia prima, recepcion, orden, producto, manufactura, liberacion y usuario responsable.",
    ])

    add_heading(doc, "4. Flujo funcional implementado", 1)
    flow_rows = [
        ("1", "Recepcion", "El operario registra materia prima recibida; se crea recepcion, inspeccion asociada, inventario de materia prima y movimiento de entrada. Se registra evidencia blockchain."),
        ("2", "Inspeccion", "La recepcion incluye checks de producto y transporte, observaciones y decision final. La decision determina si la materia prima queda aceptada para produccion."),
        ("3", "Producto/receta", "El producto define tiempos estandar, temperatura, inmersion y variantes por tamano/presentacion. Cada variante tiene materias primas requeridas por unidad."),
        ("4", "Orden de produccion", "La orden planifica fecha, codigo y productos programados. El sistema calcula insumos requeridos segun receta e inventario."),
        ("5", "Manufactura", "Se registra el hecho real: lote producido, unidades, tiempos y temperaturas reales, responsable operario y consumo de materias primas."),
        ("6", "Liberacion", "Se valida empaque/envase, factura, conductor, placa, unidades, peso y vencimiento. Si se aprueba, se crea inventario de producto terminado."),
        ("7", "Trazabilidad", "Por lote final se reconstruye origen de materias primas, inspeccion, orden, manufactura, liberacion, inventario y validaciones Fabric."),
        ("8", "Reporte/QR", "Se genera reporte PDF con estados blockchain y QR. El QR abre una vista publica de verificacion del lote."),
    ]
    add_table(doc, ["Paso", "Etapa", "Comportamiento"], flow_rows, widths=[1.0, 3.4, 13.4])

    add_heading(doc, "5. Blockchain y chaincode", 1)
    doc.add_paragraph("La integracion con Hyperledger Fabric esta separada del resto del dominio. El backend construye payloads funcionales con los datos relevantes de cada evento y llama al cliente Fabric. El chaincode calcula el hash SHA-256 dentro de Fabric, guarda la evidencia y permite validar si el payload actual coincide con lo registrado previamente.")
    add_code_paths(doc, [
        ("backend/src/modulos/blockchain/blockchain.service.js", "orquesta constructores de payload, registro y validacion de eventos criticos."),
        ("backend/src/modulos/blockchain/fabric.client.js", "cliente de conexion contra gateway/contract de Fabric."),
        ("backend/src/modulos/blockchain/fabric-traceability.service.js", "servicio especifico de comunicacion con el contrato de trazabilidad."),
        ("backend/src/modulos/blockchain/payloads/*.payload.js", "constructores de payload por evento: recepcion, inspeccion, orden, producto, manufactura, liberacion e inventarios."),
        ("fabric/chaincode/traceability/index.js", "chaincode con registrarEvento, validarEvento, consultarEvento y consultarEventosPorLote."),
        ("fabric/docker-compose.fabric.yml", "servicios Docker de CA, orderer y peers."),
        ("fabric/scripts/start.sh", "arranque de la red Fabric local."),
        ("fabric/scripts/deploy-chaincode.sh", "empaqueta, instala y despliega el chaincode."),
    ])
    add_bullets(doc, [
        "registrarEvento valida entradas, canoniza el payload, calcula hash SHA-256, guarda estado en el ledger y crea indice por lote.",
        "validarEvento reconstruye el hash del payload actual y responde VERIFICADO, ALTERADO o NO_ENCONTRADO.",
        "consultarEventosPorLote usa una composite key para recuperar eventos asociados al lote.",
        "Si Fabric no esta disponible, las pantallas pueden mostrar estados pendientes o no encontrados, porque PostgreSQL sigue siendo operativo, pero la evidencia no se confirma.",
    ])

    doc.add_page_break()
    add_heading(doc, "6. Por que se hizo asi", 1)
    why_rows = [
        ("PostgreSQL para operacion", "La trazabilidad alimentaria necesita consultas relacionales, inventarios, formularios editables y joins entre lotes, usuarios, proveedores y productos. Eso se resuelve mejor en una base relacional."),
        ("Fabric para integridad", "El ledger se usa como evidencia inmutable. Guardar solo payload/hash validable evita duplicar toda la operacion y evita simular blockchain en SQL."),
        ("Modulos por dominio", "Recepciones, produccion, liberacion, inventario y trazabilidad tienen reglas propias. Separarlos evita que un cambio en un formulario rompa toda la aplicacion."),
        ("Orden vs manufactura", "La orden es planeacion; manufactura es lo real. Esta separacion evita pedir datos reales antes de producir y permite comparar estandar contra real."),
        ("Liberacion antes de salida", "La liberacion funciona como control final de calidad y despacho interno. Solo lo aprobado alimenta inventario de producto terminado."),
        ("Payloads normalizados", "Los hashes deben ser reproducibles. Por eso se excluyen IDs tecnicos y metadata variable, y se ordenan campos antes de enviar a Fabric."),
        ("QR publico", "El QR no reemplaza el reporte: sirve para verificacion rapida del lote desde fuera del sistema privado, sin exponer recetas ni datos sensibles."),
        ("JWT y cierre por inactividad", "El sistema maneja informacion operativa interna. La sesion con token y expiracion por actividad reduce riesgo de dejar cuentas abiertas."),
    ]
    add_table(doc, ["Decision", "Justificacion"], why_rows, widths=[4.0, 13.8])

    add_heading(doc, "7. Pendientes recomendados para cierre", 1)
    add_bullets(doc, [
        "Crear modulo de almacenamiento si el alcance final exige temperatura, ubicacion y fechas de entrada/salida por lote.",
        "Decidir si liberacion reemplaza despacho o si se implementara un despacho independiente con cliente receptor y temperatura de transporte.",
        "Agregar modulo de saneamiento para RF15 y modulo de devoluciones/no conformidades para RF16.",
        "Completar reportes con exportacion Excel y opcion de varios lotes.",
        "Separar informacion publica de consumidor e informacion para INVIMA dentro del portal QR.",
        "Formalizar pruebas de rendimiento, usabilidad y disponibilidad para los requerimientos no funcionales.",
        "Documentar despliegue Fabric productivo si se pasa de Docker local a servidores fisicos o nube.",
    ])

    add_heading(doc, "8. Conclusion", 1)
    doc.add_paragraph("Trazaap ya implementa el nucleo funcional de trazabilidad: recepcion, inspeccion, inventario de materias primas, productos/recetas, ordenes de produccion, manufactura real, liberacion, inventario de producto terminado, trazabilidad por lote, QR, reportes y validacion blockchain con Hyperledger Fabric. Lo pendiente se concentra en modulos complementarios del alcance normativo, pruebas no funcionales formales y algunas extensiones de despacho, almacenamiento, saneamiento y devoluciones.")

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Trazaap - documento de soporte tecnico y funcional")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("64748B")

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build_doc()
    print(path)
